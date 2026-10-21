"""One visual style for every workbook sheet: header, borders, links, fills, print setup."""

from __future__ import annotations

from collections.abc import Sequence

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from tiktok_trip.export.palette import readable_text_color

MAX_FORMULA_TEXT = 250

HEADER_FILL = PatternFill("solid", fgColor="37474F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(vertical="center", wrap_text=True)
LINK_FONT = Font(color="0563C1", underline="single")
DAY_FILL = PatternFill("solid", fgColor="E0E0E0")
THIN = Side(style="thin", color="B0B7BB")
MEDIUM = Side(style="medium", color="455A64")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
NUMBER_TOP = Alignment(horizontal="right", vertical="top")


def fill_for(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.lstrip("#"))


def text_rgb(background_hex: str) -> str:
    return readable_text_color(background_hex).lstrip("#")


def set_link(cell: Cell | MergedCell, url: str, label: str) -> None:
    """HYPERLINK() formula, or a plain hyperlink when the URL or label is too long for one."""
    if len(url) > MAX_FORMULA_TEXT or len(label) > MAX_FORMULA_TEXT:
        cell.value = label
        cell.hyperlink = url
    else:
        safe_url = url.replace('"', '""')
        safe_label = label.replace('"', '""')
        cell.value = f'=HYPERLINK("{safe_url}","{safe_label}")'
    cell.font = LINK_FONT


def style_header_cell(cell: Cell | MergedCell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = HEADER_ALIGN
    cell.border = GRID


def write_header(
    ws: Worksheet,
    columns: Sequence[str],
    widths: Sequence[int] | None = None,
    *,
    row: int = 1,
    first_col: int = 1,
) -> None:
    for offset, name in enumerate(columns):
        cell = ws.cell(row=row, column=first_col + offset, value=name)
        style_header_cell(cell)
        width = widths[offset] if widths and offset < len(widths) else 14
        ws.column_dimensions[get_column_letter(first_col + offset)].width = width


def style_body(ws: Worksheet, first_row: int, last_row: int, first_col: int, last_col: int) -> None:
    """Thin grid, wrapped top-aligned text, numbers right-aligned; keeps fonts and fills."""
    for row in ws.iter_rows(
        min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col
    ):
        for cell in row:
            cell.border = GRID
            numeric = isinstance(cell.value, int | float) and not isinstance(cell.value, bool)
            cell.alignment = NUMBER_TOP if numeric else WRAP_TOP


def outline(ws: Worksheet, first_row: int, last_row: int, first_col: int, last_col: int) -> None:
    """Medium frame around a block, thin lines inside."""
    for row in range(first_row, last_row + 1):
        for col in range(first_col, last_col + 1):
            ws.cell(row=row, column=col).border = Border(
                left=MEDIUM if col == first_col else THIN,
                right=MEDIUM if col == last_col else THIN,
                top=MEDIUM if row == first_row else THIN,
                bottom=MEDIUM if row == last_row else THIN,
            )


def print_setup(ws: Worksheet, *, title_rows: str | None = "1:1") -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if title_rows:
        ws.print_title_rows = title_rows
