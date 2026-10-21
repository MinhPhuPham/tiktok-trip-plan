"""Plan-to-go sheet: a one-row-per-day summary, every stop day by day, then unplaced places."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tiktok_trip.clock import to_minutes
from tiktok_trip.config import PlanConfig
from tiktok_trip.export.dataset import ExportData, first_sentence, price_text
from tiktok_trip.export.excel_style import (
    fill_for,
    print_setup,
    set_link,
    style_body,
    style_header_cell,
    text_rgb,
    write_header,
)
from tiktok_trip.export.i18n import (
    Labels,
    best_with_why,
    center_name,
    clock,
    clock_range,
    kind_label,
    opening_hours,
    stop_reason,
    tags_text,
    unplaced_text,
)
from tiktok_trip.export.palette import day_color
from tiktok_trip.geo import coord_maps_url
from tiktok_trip.models import DayPlan, NearbySpot, Place, RouteStop

PLAN_WIDTHS = (8, 5, 24, 34, 10, 13, 32, 10, 10, 11, 12, 28, 30, 20, 16, 56)
LEG_WIDTH = 12
MINUTES_FORMAT = "0"
TITLE_FONT = Font(bold=True, size=14, color="263238")
MUTED_FONT = Font(italic=True, color="607D8B")
NO_WRAP = Alignment(vertical="center", wrap_text=False)
TITLE_ROW = 1
SUMMARY_HEADER_ROW = 2
SUMMARY_COLS = (1, 2, 5, 6, 7, 8, 9, 10)
"""Day, Areas (merged over B:D), Stops, Start, End, Total, Travel, Directions (one per leg)."""
AREAS_LAST_COL = 4
SUMMARY_MINUTE_COLS = (8, 9)
PLACE_COL = 4
PRIORITY_COL = 5
AREA_COL = 7
NOTE_COL = 16
MINUTE_COLS = (10, 11)
VISIT_KINDS = frozenset({"place", "nearby"})
CENTER_KINDS = frozenset({"start", "end", "rest"})
TIMED_KINDS = frozenset({"place", "nearby", "rest"})
AGENT_GROUPING = "agent"


@dataclass(frozen=True)
class Lookups:
    places: dict[str, Place]
    spots: dict[str, NearbySpot]
    zone_sectors: dict[str, str]

    @classmethod
    def build(cls, data: ExportData) -> Lookups:
        zones = data.spatial.zones if data.spatial else []
        return cls(
            places={place.id: place for place in data.places},
            spots={spot.id: spot for spot in data.nearby},
            zone_sectors={zone.id: zone.sector for zone in zones},
        )

    def sector(self, zone_id: str) -> str:
        return self.zone_sectors.get(zone_id) or zone_id.split("-", maxsplit=1)[0]

    def place(self, stop: RouteStop) -> Place | None:
        return self.places.get(stop.ref_id or "") if stop.kind == "place" else None

    def spot(self, stop: RouteStop) -> NearbySpot | None:
        return self.spots.get(stop.ref_id or "") if stop.kind == "nearby" else None


@dataclass(frozen=True)
class StopRow:
    values: list[object]
    url: str
    label: str


def days_of(data: ExportData) -> list[DayPlan]:
    return data.itinerary.days if data.itinerary else []


def areas_text(day: DayPlan, lookups: Lookups) -> str:
    """'S, SW · S-1, SW-1': sectors in visiting order, then the zones."""
    sectors = list(dict.fromkeys(lookups.sector(zone) for zone in day.zone_ids))
    return " · ".join(part for part in (", ".join(sectors), ", ".join(day.zone_ids)) if part)


def area_text(labels: Labels, sector: str | None, zone: str | None) -> str:
    text = labels.direction(sector)
    return f"{text} · {labels.zone} {zone}" if zone else text


def travel_minutes(day: DayPlan) -> int:
    return sum(stop.travel_min_to_next or 0 for stop in day.stops)


def visit_count(day: DayPlan) -> int:
    return sum(stop.kind in VISIT_KINDS for stop in day.stops)


def day_heading(labels: Labels, day: DayPlan, lookups: Lookups) -> str:
    """'Day 1 · N, NE · N-1, NE-1 · 8:00 AM-8:05 PM' (en dash in the sheet)."""
    parts = [labels.day.format(day=day.day), areas_text(day, lookups)]
    if day.stops:
        parts.append(clock_range(labels, day.stops[0].depart, day.stops[-1].arrive))
    return " · ".join(part for part in parts if part)


def plan_title(labels: Labels, data: ExportData) -> str:
    itinerary = data.itinerary
    grouping = itinerary.grouping if itinerary else data.config.trip.day_grouping
    mode = itinerary.travel_mode if itinerary else data.config.travel.mode
    method = (
        labels.agent_method
        if grouping == AGENT_GROUPING
        else labels.solver_method.format(grouping=labels.groupings.get(grouping, grouping))
    )
    return labels.plan_title.format(
        city=data.config.city,
        days=labels.day_count(len(itinerary.days) if itinerary else 0),
        method=method,
        mode=labels.travel_modes.get(mode, mode),
    )


def activity_text(labels: Labels, plan: PlanConfig, slot_id: str | None) -> str:
    """The rhythm slot's name in the sheet's language; the raw id when it is not configured."""
    if not slot_id:
        return ""
    slot = plan.slot(slot_id)
    if slot is None:
        return slot_id
    return slot.name_vi if labels.code == "vi" else slot.name_en


def rest_label(labels: Labels, data: ExportData) -> str:
    center = data.spatial.center if data.spatial else None
    return labels.rest_at_hotel if center and center.source == "hotel" else labels.rest_at_center


def stop_target(
    labels: Labels, stop: RouteStop, data: ExportData, lookups: Lookups
) -> tuple[str, str]:
    """Maps URL and link label of a stop: the center for start, end and rest stops."""
    if stop.kind in CENTER_KINDS:
        center = data.spatial.center if data.spatial else None
        if center is not None:
            return center.maps_url, center_name(labels, center)
    elif place := lookups.place(stop):
        return place.maps_url, place.name
    elif spot := lookups.spot(stop):
        return spot.maps_url, spot.name
    return coord_maps_url(stop.lat, stop.lng), stop.name


def stop_row(
    labels: Labels,
    stop: RouteStop,
    *,
    day: int,
    number: int | None,
    data: ExportData,
    lookups: Lookups,
) -> StopRow:
    """One detail row; the Place cell stays empty here and is filled with a Maps link."""
    url, label = stop_target(labels, stop, data, lookups)
    activity = activity_text(labels, data.config.plan, stop.slot)
    area, priority, hours, best, tags, price, note = "", "", "", "", "", "", ""
    if stop.kind in CENTER_KINDS:
        area = labels.center_area
        if stop.kind == "rest":
            label = rest_label(labels, data)
    elif place := lookups.place(stop):
        sector, zone = (place.spatial.sector, place.spatial.zone) if place.spatial else (None, None)
        area = area_text(labels, sector, zone)
        priority = labels.priority(place.priority)
        hours = opening_hours(labels, place.hours)
        best = best_with_why(labels, place.best_time)
        tags = tags_text(labels, place.tags)
        price = price_text(place)
        note = first_sentence(place.notes)
    elif spot := lookups.spot(stop):
        area = area_text(labels, lookups.sector(spot.zone_id), spot.zone_id)
        activity = kind_label(labels, spot.kind)
    reason = stop_reason(labels, stop.reason, stop.reason_vi)
    note = " · ".join(part for part in (reason, note) if part)
    values: list[object] = [
        day,
        number,
        activity or None,
        None,
        priority or None,
        labels.stop_kinds[stop.kind],
        area,
        clock(labels, stop.arrive) if stop.kind != "start" else None,
        clock(labels, stop.depart) if stop.kind != "end" else None,
        to_minutes(stop.depart) - to_minutes(stop.arrive) if stop.kind in TIMED_KINDS else None,
        stop.travel_min_to_next,
        hours or None,
        best or None,
        tags or None,
        price or None,
        note or None,
    ]
    return StopRow(values=values, url=url, label=label)


def day_rows(labels: Labels, day: DayPlan, data: ExportData, lookups: Lookups) -> list[StopRow]:
    rows: list[StopRow] = []
    number = 0
    for stop in day.stops:
        visit = stop.kind in VISIT_KINDS
        number += visit
        rows.append(
            stop_row(
                labels,
                stop,
                day=day.day,
                number=number if visit else None,
                data=data,
                lookups=lookups,
            )
        )
    return rows


def _style_heading(ws: Worksheet, row: int, day: int, last_col: int) -> None:
    color = day_color(day)
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill_for(color)
        cell.font = Font(bold=True, color=text_rgb(color))
        cell.alignment = NO_WRAP


def _direction_links(ws: Worksheet, row: int, first_col: int, labels: Labels, day: DayPlan) -> None:
    single = len(day.maps_urls) == 1
    for leg, url in enumerate(day.maps_urls, start=1):
        label = labels.directions_link if single else labels.leg.format(leg=leg)
        set_link(ws.cell(row=row, column=first_col + leg - 1), url, label)


def _write_summary(ws: Worksheet, data: ExportData, labels: Labels, lookups: Lookups) -> int:
    """Title, then one row per day under its own header; returns the last row used."""
    ws.cell(row=TITLE_ROW, column=1, value=plan_title(labels, data)).font = TITLE_FONT
    ws.row_dimensions[TITLE_ROW].height = 22
    days = days_of(data)
    legs = max((len(day.maps_urls) for day in days), default=1)
    links_col = SUMMARY_COLS[-1]
    last_col = links_col + max(legs, 1) - 1
    for col, name in zip(SUMMARY_COLS, labels.summary_columns, strict=True):
        ws.cell(row=SUMMARY_HEADER_ROW, column=col, value=name)
    for col in range(1, last_col + 1):
        style_header_cell(ws.cell(row=SUMMARY_HEADER_ROW, column=col))
    ws.merge_cells(
        start_row=SUMMARY_HEADER_ROW,
        start_column=2,
        end_row=SUMMARY_HEADER_ROW,
        end_column=AREAS_LAST_COL,
    )
    for extra in range(len(PLAN_WIDTHS) + 1, last_col + 1):
        ws.column_dimensions[get_column_letter(extra)].width = LEG_WIDTH
    if not days:
        ws.cell(row=SUMMARY_HEADER_ROW + 1, column=1, value=labels.no_plan).font = MUTED_FONT
        return SUMMARY_HEADER_ROW + 1
    row = SUMMARY_HEADER_ROW
    for day in days:
        row += 1
        values: list[object] = [
            labels.day.format(day=day.day),
            areas_text(day, lookups),
            visit_count(day),
            clock(labels, day.stops[0].depart) if day.stops else None,
            clock(labels, day.stops[-1].arrive) if day.stops else None,
            day.total_minutes,
            travel_minutes(day),
        ]
        for col, value in zip(SUMMARY_COLS, values, strict=False):
            ws.cell(row=row, column=col, value=value)
        _direction_links(ws, row, links_col, labels, day)
    style_body(ws, SUMMARY_HEADER_ROW + 1, row, 1, last_col)
    for day_row, day in enumerate(days, start=SUMMARY_HEADER_ROW + 1):
        ws.merge_cells(
            start_row=day_row, start_column=2, end_row=day_row, end_column=AREAS_LAST_COL
        )
        color = day_color(day.day)
        cell = ws.cell(row=day_row, column=1)
        cell.fill = fill_for(color)
        cell.font = Font(bold=True, color=text_rgb(color))
        for col in SUMMARY_MINUTE_COLS:
            ws.cell(row=day_row, column=col).number_format = MINUTES_FORMAT
    return row


def _write_stops(ws: Worksheet, first_row: int, rows: Sequence[StopRow]) -> int:
    row = first_row - 1
    for stop in rows:
        row += 1
        for col, value in enumerate(stop.values, start=1):
            ws.cell(row=row, column=col, value=value)
        set_link(ws.cell(row=row, column=PLACE_COL), stop.url, stop.label)
    return row


def _write_detail(
    ws: Worksheet, header_row: int, data: ExportData, labels: Labels, lookups: Lookups
) -> int:
    """Detail table: header, then a coloured heading and the stop rows per day; returns last row."""
    columns = labels.plan_columns
    write_header(ws, columns, PLAN_WIDTHS, row=header_row)
    row = header_row
    headings: list[tuple[int, int]] = []
    for day in days_of(data):
        row += 1
        ws.cell(row=row, column=1, value=day_heading(labels, day, lookups))
        headings.append((row, day.day))
        row = _write_stops(ws, row + 1, day_rows(labels, day, data, lookups))
    if row > header_row:
        style_body(ws, header_row + 1, row, 1, len(columns))
        for cell_row in range(header_row + 1, row + 1):
            for col in MINUTE_COLS:
                ws.cell(row=cell_row, column=col).number_format = MINUTES_FORMAT
    for heading_row, day_number in headings:
        _style_heading(ws, heading_row, day_number, len(columns))
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{row}"
    return row


def _write_unplaced(
    ws: Worksheet, first_row: int, data: ExportData, labels: Labels, lookups: Lookups
) -> int:
    """Places the plan left out, each with its reason; returns the last row used."""
    unplaced = data.itinerary.unplaced if data.itinerary else []
    if not unplaced:
        return first_row - 1
    ws.cell(row=first_row, column=1, value=labels.not_planned).font = Font(bold=True)
    row = first_row
    for item in unplaced:
        row += 1
        place = lookups.places.get(item.ref_id)
        target = ws.cell(row=row, column=PLACE_COL)
        if place is None:
            target.value = item.ref_id
        else:
            set_link(target, place.maps_url, place.name)
            ws.cell(row=row, column=PRIORITY_COL, value=labels.priority(place.priority) or None)
            sector, zone = (
                (place.spatial.sector, place.spatial.zone) if place.spatial else (None, None)
            )
            ws.cell(row=row, column=AREA_COL, value=area_text(labels, sector, zone))
        ws.cell(row=row, column=NOTE_COL, value=unplaced_text(labels, item.code, item.reason))
    style_body(ws, first_row + 1, row, PLACE_COL, NOTE_COL)
    return row


def build_plan(ws: Worksheet, data: ExportData, labels: Labels) -> None:
    lookups = Lookups.build(data)
    summary_end = _write_summary(ws, data, labels, lookups)
    header_row = summary_end + 2
    detail_end = _write_detail(ws, header_row, data, labels, lookups)
    _write_unplaced(ws, detail_end + 2, data, labels, lookups)
    print_setup(ws, title_rows=f"{header_row}:{header_row}")
