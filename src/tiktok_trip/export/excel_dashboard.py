"""Dashboard sheet: title, totals, 3x3 compass of linked place grids, plan at a glance beside it."""

from __future__ import annotations

import math
from collections.abc import Sequence
from dataclasses import dataclass, field

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tiktok_trip.export.dataset import (
    ExportData,
    by_priority,
    places_by_sector,
    sort_places,
    starred_name,
    truncate,
    unique_videos,
)
from tiktok_trip.export.excel_plan import (
    Lookups,
    activity_text,
    day_heading,
    days_of,
    stop_target,
)
from tiktok_trip.export.excel_style import (
    LINK_FONT,
    WRAP_TOP,
    fill_for,
    outline,
    print_setup,
    set_link,
    style_body,
    text_rgb,
    write_header,
)
from tiktok_trip.export.i18n import (
    Labels,
    center_name,
    clock,
    clock_range,
    hours_and_best,
    kind_label,
    stop_reason,
    tags_text,
)
from tiktok_trip.export.palette import CENTER_COLOR, UNASSIGNED, day_color, sector_color
from tiktok_trip.models import Place, RouteStop

TITLE_ROW = 1
TOTALS_ROW = 2
COMPASS_TOP = 4
FIRST_COL = 2
BLOCK_COLS = 2
MIN_BLOCK_PLACES = 3
COMPASS_WIDTHS = (28, 32)
GLANCE_WIDTHS = (24, 32, 17, 42)
GLANCE_NOTE_CHARS = 60
LINE_HEIGHT = 15.0
MAX_ROW_HEIGHT = 45.0
MUTED_FONT = Font(italic=True, color="607D8B")

COMPASS_LAYOUT_8 = (("NW", "N", "NE"), ("W", None, "E"), ("SW", "S", "SE"))
COMPASS_LAYOUT_4 = ((None, "N", None), ("W", None, "E"), (None, "S", None))
CENTER_SLOT = (1, 1)


def compass_layout(sector_count: int) -> tuple[tuple[str | None, ...], ...]:
    return COMPASS_LAYOUT_4 if sector_count == 4 else COMPASS_LAYOUT_8


def block_height(data: ExportData) -> int:
    """Place rows per compass block: the busiest sector shown, never fewer than three."""
    grouped = places_by_sector(data.places)
    shown = [name for line in compass_layout(data.sector_count) for name in line if name]
    return max([MIN_BLOCK_PLACES, *(len(grouped.get(name, [])) for name in shown)])


def glance_first_col() -> int:
    """First column of the plan at a glance: right of the compass and one spacer column."""
    return FIRST_COL + 3 * BLOCK_COLS + 1


def block_origin(grid_row: int, grid_col: int, height: int) -> tuple[int, int]:
    return COMPASS_TOP + grid_row * (height + 1), FIRST_COL + grid_col * BLOCK_COLS


@dataclass
class RowHeights:
    """Estimated wrapped line count per row, so shared compass/table rows are tall enough."""

    widths: dict[int, float] = field(default_factory=dict)
    lines: dict[int, int] = field(default_factory=dict)

    def note(self, row: int, col: int, text: str) -> None:
        width = max(self.widths.get(col, 10.0) - 1, 1.0)
        count = sum(max(1, math.ceil(len(line) * 1.1 / width)) for line in text.split("\n"))
        self.lines[row] = max(self.lines.get(row, 1), count)

    def apply(self, ws: Worksheet) -> None:
        for row, count in self.lines.items():
            if count > 1:
                ws.row_dimensions[row].height = min(count * LINE_HEIGHT, MAX_ROW_HEIGHT)


def _write_title(ws: Worksheet, data: ExportData, labels: Labels) -> None:
    date = data.generated_at.strftime(labels.date_format)
    cell = ws.cell(row=TITLE_ROW, column=FIRST_COL)
    cell.value = labels.title.format(city=data.config.city, date=date)
    cell.font = Font(bold=True, size=16, color="263238")
    ws.row_dimensions[TITLE_ROW].height = 24


def _write_totals(ws: Worksheet, data: ExportData, labels: Labels) -> None:
    unassigned = len(places_by_sector(data.places).get(UNASSIGNED, []))
    totals: list[tuple[str, int]] = [
        ("places", len(data.places)),
        ("zones", len(data.spatial.zones) if data.spatial else 0),
        ("videos", len(unique_videos(data.places))),
        ("pending_review", data.pending_review),
        ("high_priority", sum(place.priority == "high" for place in data.places)),
    ]
    if must := sum(place.priority == "must" for place in data.places):
        totals.append(("must_go", must))
    if unassigned:
        totals.insert(1, ("without_direction", unassigned))
    last_col = FIRST_COL + 3 * BLOCK_COLS - 1
    ws.merge_cells(
        start_row=TOTALS_ROW, start_column=FIRST_COL, end_row=TOTALS_ROW, end_column=last_col
    )
    cell = ws.cell(row=TOTALS_ROW, column=FIRST_COL)
    cell.value = " · ".join(f"{labels.totals[key]} {value}" for key, value in totals)
    cell.font = Font(bold=True, color="455A64")
    cell.alignment = Alignment(horizontal="left")


def _block_header(ws: Worksheet, row: int, col: int, text: str, hex_color: str) -> None:
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + BLOCK_COLS - 1)
    cell = ws.cell(row=row, column=col, value=text)
    cell.fill = fill_for(hex_color)
    cell.font = Font(bold=True, size=12, color=text_rgb(hex_color))
    cell.alignment = Alignment(horizontal="center", vertical="center")


def compass_times(labels: Labels, place: Place) -> str:
    """Hours with the best time, then the tag chips on their own line."""
    times = hours_and_best(labels, place.hours, place.best_time)
    tags = tags_text(labels, place.tags)
    return f"{times}\n{tags}" if tags else times


def _sector_block(
    ws: Worksheet,
    *,
    origin: tuple[int, int],
    height: int,
    sector: str,
    places: Sequence[Place],
    data: ExportData,
    labels: Labels,
    heights: RowHeights,
) -> None:
    row, col = origin
    _block_header(ws, row, col, f"{sector} ({len(places)})", sector_color(sector))
    ordered = by_priority(sort_places(places, data.itinerary))
    for place_row, place in enumerate(ordered, start=row + 1):
        name = starred_name(place)
        set_link(ws.cell(row=place_row, column=col), place.maps_url, name)
        times = compass_times(labels, place)
        ws.cell(row=place_row, column=col + 1, value=times)
        for offset, text in enumerate((name, times)):
            ws.cell(row=place_row, column=col + offset).alignment = WRAP_TOP
            heights.note(place_row, col + offset, text)
    outline(ws, row, row + height, col, col + BLOCK_COLS - 1)


def _center_block(
    ws: Worksheet, origin: tuple[int, int], height: int, data: ExportData, labels: Labels
) -> None:
    row, col = origin
    _block_header(ws, row, col, labels.center, CENTER_COLOR)
    last_col = col + BLOCK_COLS - 1
    for offset in range(1, 4):
        ws.merge_cells(
            start_row=row + offset, start_column=col, end_row=row + offset, end_column=last_col
        )
        ws.cell(row=row + offset, column=col).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    center = data.spatial.center if data.spatial else None
    if center is None:
        ws.cell(row=row + 1, column=col, value=labels.center_missing)
    else:
        name = ws.cell(row=row + 1, column=col)
        set_link(name, center.maps_url, center_name(labels, center))
        name.font = Font(bold=True, size=12, color=LINK_FONT.color, underline="single")
        source = labels.center_sources.get(center.source, center.source)
        ws.cell(row=row + 2, column=col, value=f"{labels.source}: {source}")
        coords = ws.cell(row=row + 3, column=col, value=f"{center.lat:.5f}, {center.lng:.5f}")
        coords.font = Font(color="607D8B")
    outline(ws, row, row + height, col, last_col)


def _write_compass(ws: Worksheet, data: ExportData, labels: Labels, heights: RowHeights) -> int:
    """Draws the compass and returns its last row."""
    height = block_height(data)
    grouped = places_by_sector(data.places)
    for grid_row, line in enumerate(compass_layout(data.sector_count)):
        for grid_col, sector in enumerate(line):
            origin = block_origin(grid_row, grid_col, height)
            if sector is not None:
                _sector_block(
                    ws,
                    origin=origin,
                    height=height,
                    sector=sector,
                    places=grouped.get(sector, []),
                    data=data,
                    labels=labels,
                    heights=heights,
                )
            elif (grid_row, grid_col) == CENTER_SLOT:
                _center_block(ws, origin, height, data, labels)
    return COMPASS_TOP + 3 * (height + 1) - 1


def glance_activity(labels: Labels, stop: RouteStop, data: ExportData, lookups: Lookups) -> str:
    """Slot name in the sheet's language; a nearby spot's kind ('Temple / pagoda'); without a slot
    the category or the stop type."""
    if spot := lookups.spot(stop):
        return kind_label(labels, spot.kind)
    slot = activity_text(labels, data.config.plan, stop.slot)
    if slot:
        return slot
    place = lookups.place(stop)
    return labels.categories[place.category] if place else labels.stop_kinds[stop.kind]


def glance_time(labels: Labels, stop: RouteStop) -> str:
    """'8:12-8:57 AM' (or '8:12-8:57'); a single clock for the start and end of the day."""
    if stop.kind == "start":
        return clock(labels, stop.depart)
    if stop.kind == "end":
        return clock(labels, stop.arrive)
    return clock_range(labels, stop.arrive, stop.depart)


def _day_header(ws: Worksheet, row: int, columns: tuple[int, int], text: str, day: int) -> None:
    first_col, last_col = columns
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    color = day_color(day)
    cell = ws.cell(row=row, column=first_col, value=text)
    cell.fill = fill_for(color)
    cell.font = Font(bold=True, color=text_rgb(color))
    cell.alignment = Alignment(vertical="center")


def _write_glance(ws: Worksheet, data: ExportData, labels: Labels, heights: RowHeights) -> int:
    """One block per day right of the compass: a coloured header, then one row per stop."""
    first_col = glance_first_col()
    last_col = first_col + len(GLANCE_WIDTHS) - 1
    days = days_of(data)
    if not days:
        ws.cell(row=COMPASS_TOP, column=first_col, value=labels.no_plan).font = MUTED_FONT
        return COMPASS_TOP
    write_header(ws, labels.glance_columns, GLANCE_WIDTHS, row=COMPASS_TOP, first_col=first_col)
    lookups = Lookups.build(data)
    row = COMPASS_TOP
    for day in days:
        row += 1
        top = row
        _day_header(ws, row, (first_col, last_col), day_heading(labels, day, lookups), day.day)
        for stop in day.stops:
            row += 1
            url, label = stop_target(labels, stop, data, lookups)
            texts = (
                glance_activity(labels, stop, data, lookups),
                label,
                glance_time(labels, stop),
                truncate(stop_reason(labels, stop.reason, stop.reason_vi), GLANCE_NOTE_CHARS),
            )
            for offset, text in enumerate(texts):
                ws.cell(row=row, column=first_col + offset, value=text or None)
                heights.note(row, first_col + offset, text)
            style_body(ws, row, row, first_col, last_col)
            set_link(ws.cell(row=row, column=first_col + 1), url, label)
        outline(ws, top, row, first_col, last_col)
    return row


def _write_legend(ws: Worksheet, row: int, data: ExportData, labels: Labels) -> int:
    has_must = any(place.priority == "must" for place in data.places)
    legend = labels.priority_legend(include_must=has_must)
    ws.cell(row=row, column=FIRST_COL, value=legend).font = MUTED_FONT
    return row


def _write_warnings(ws: Worksheet, data: ExportData, labels: Labels, first_row: int) -> int:
    row = first_row
    for warning in data.spatial.warnings if data.spatial else []:
        cell = ws.cell(row=row, column=FIRST_COL, value=f"{labels.warning}: {warning}")
        cell.font = Font(color="BF360C", italic=True)
        row += 1
    return row - 1


def build_dashboard(ws: Worksheet, data: ExportData, labels: Labels) -> None:
    ws.column_dimensions["A"].width = 2
    widths: dict[int, float] = {}
    for grid_col in range(3):
        for offset, width in enumerate(COMPASS_WIDTHS):
            col = FIRST_COL + grid_col * BLOCK_COLS + offset
            widths[col] = width
            ws.column_dimensions[get_column_letter(col)].width = width
    ws.column_dimensions[get_column_letter(glance_first_col() - 1)].width = 3
    for offset, width in enumerate(GLANCE_WIDTHS):
        col = glance_first_col() + offset
        widths[col] = width
        ws.column_dimensions[get_column_letter(col)].width = width
    heights = RowHeights(widths=widths)

    _write_title(ws, data, labels)
    _write_totals(ws, data, labels)
    compass_end = _write_compass(ws, data, labels, heights)
    legend_end = _write_legend(ws, compass_end + 1, data, labels)
    glance_end = _write_glance(ws, data, labels, heights)
    warnings_end = _write_warnings(ws, data, labels, legend_end + 1)
    heights.apply(ws)

    last_col = glance_first_col() + len(GLANCE_WIDTHS) - 1
    last_row = max(legend_end, glance_end, warnings_end)
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    print_setup(ws, title_rows=None)
    ws.sheet_view.zoomScale = 90
