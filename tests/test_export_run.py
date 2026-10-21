from __future__ import annotations

import dataclasses
import json
import logging
import re

import httpx
import pytest
import respx
from fixtures.sample_dataset import build_dataset
from openpyxl import load_workbook

from tiktok_trip import export
from tiktok_trip.config import Secrets
from tiktok_trip.context import Context
from tiktok_trip.errors import UserFacingError
from tiktok_trip.freshness import mark_enriched
from tiktok_trip.http import RateLimiter
from tiktok_trip.models import Candidate, Place, ReviewRow
from tiktok_trip.storage import load_models, save_model, save_models, write_csv


def seed(ctx: Context) -> None:
    data = build_dataset()
    paths = ctx.paths
    save_models(paths.places_file, data.places)
    assert data.spatial is not None
    save_model(paths.spatial_file, data.spatial)
    save_models(paths.nearby_file, data.nearby)
    assert data.itinerary is not None
    save_model(paths.routes_file, data.itinerary)
    pending = ReviewRow(
        candidate=Candidate(
            video_id="v1", video_url="https://tiktok.com/v1", name="Mystery", confidence=0.4
        ),
        reason="not_found",
    )
    save_models(paths.review_file, [pending])
    write_csv(paths.review_csv, ["key"], [[pending.key]])


def test_run_writes_layout(ctx: Context) -> None:
    seed(ctx)
    export.run(ctx)
    out = ctx.paths.output
    for name in (
        "map.html",
        "app/index.html",
        "app/trip.json",
        "app/offline.json",
        "places.json",
        "review.csv",
        "mymaps/mymaps_N.csv",
        "trips/en/trip.xlsx",
        "trips/vi/trip.xlsx",
    ):
        assert (out / name).exists(), name
    assert ctx.paths.trip_xlsx("vi") == out / "trips" / "vi" / "trip.xlsx"
    assert sorted(p.relative_to(out).as_posix() for p in out.rglob("*.xlsx")) == [
        "trips/en/trip.xlsx",
        "trips/vi/trip.xlsx",
    ]
    assert (out / "places.json").read_bytes() == ctx.paths.places_file.read_bytes()
    assert not list(out.glob("mymaps_*.csv"))
    en = load_workbook(ctx.paths.trip_xlsx("en"))
    assert en.sheetnames == ["Dashboard", "Plan-to-go"]
    dashboard = en["Dashboard"]
    assert "Pending review 1" in str(dashboard["B2"].value)
    vi = load_workbook(ctx.paths.trip_xlsx("vi"))
    assert vi.sheetnames == ["Tổng Hợp", "Lịch Trình"]
    assert str(vi["Tổng Hợp"]["B2"].value).startswith("Địa điểm")


def test_run_removes_stale_outputs(ctx: Context) -> None:
    seed(ctx)
    out = ctx.paths.output
    stale = [
        out / "trip.xlsx",
        out / "mymaps_N.csv",
        out / "mymaps_NW.csv",
        *(
            out / "trips" / lang / name
            for lang in ("en", "vi")
            for name in ("dashboard.xlsx", "plan-to-go.xlsx")
        ),
    ]
    for path in stale:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("old")
    (ctx.paths.mymaps_dir).mkdir(parents=True)
    (ctx.paths.mymaps_dir / "mymaps_NW.csv").write_text("old")
    export.run(ctx)
    assert not any(path.exists() for path in stale)
    assert not (ctx.paths.mymaps_dir / "mymaps_NW.csv").exists()
    assert (ctx.paths.mymaps_dir / "mymaps_N.csv").exists()
    assert ctx.paths.trip_xlsx("en").exists()


def test_run_without_places_is_a_noop(ctx: Context) -> None:
    export.run(ctx)
    assert not ctx.paths.trips_dir.exists()


def test_run_without_spatial_or_routes(ctx: Context) -> None:
    data = build_dataset()
    save_models(ctx.paths.places_file, data.places)
    export.run(ctx)
    workbook = load_workbook(ctx.paths.trip_xlsx("en"))
    assert workbook.sheetnames == ["Dashboard", "Plan-to-go"]
    assert workbook["Dashboard"]["I4"].value == "Plan not generated yet: run tiktok-trip route"
    assert ctx.paths.map_html.exists()


@respx.mock
def test_run_draws_routes_along_roads_from_the_cache_on_rerun(
    ctx: Context, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(RateLimiter, "wait", lambda self: None)
    seed(ctx)
    ctx.config.roads.enabled = True
    route = respx.get(
        re.compile(r"https://routing\.openstreetmap\.de/routed-car/route/v1/.*")
    ).mock(
        return_value=httpx.Response(
            200,
            json={
                "code": "Ok",
                "routes": [{"geometry": "_p~iF~ps|U", "distance": 1, "duration": 1}],
            },
        )
    )
    export.run(ctx)
    legs = route.call_count
    assert legs > 0
    export.run(ctx)
    assert route.call_count == legs
    html = ctx.paths.map_html.read_text(encoding="utf-8")
    payload = json.loads(
        re.search(r'id="trip-data" type="application/json">(.*?)</script>', html, re.DOTALL)[1]
    )  # type: ignore[index]
    assert len(payload["roads"]) == legs
    assert payload["roads"]["center>gp:pho1"] == "_p~iF~ps|U"


def test_run_fails_when_server_key_is_in_output(ctx: Context) -> None:
    seed(ctx)
    keyed = dataclasses.replace(ctx, secrets=Secrets(google_maps_key="AIzaSERVER-secret-0123"))
    export.run(keyed)
    (ctx.paths.output / "extra.txt").write_text("key=AIzaSERVER-secret-0123")
    with pytest.raises(UserFacingError, match=r"extra\.txt"):
        export.run(keyed)


def test_run_warns_when_places_were_not_enriched(
    ctx: Context, caplog: pytest.LogCaptureFixture
) -> None:
    seed(ctx)
    with caplog.at_level(logging.WARNING):
        export.run(ctx)
    assert any("tiktok-trip enrich" in r.getMessage() for r in caplog.records)
    caplog.clear()
    mark_enriched(ctx.paths, load_models(ctx.paths.places_file, Place))
    with caplog.at_level(logging.WARNING):
        export.run(ctx)
    assert not any("enrich" in r.getMessage() for r in caplog.records)
