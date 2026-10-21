"""Readable ink on sector fills, hours across midnight, translated kinds, months and plan notes."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from fixtures.sample_dataset import build_dataset
from openpyxl import load_workbook

from tiktok_trip.clock import EN_DASH
from tiktok_trip.export.excel import write_trip
from tiktok_trip.export.i18n import (
    EN,
    VI,
    clock_range,
    kind_label,
    localize_calendar,
    opening_hours,
    stop_reason,
    unplaced_text,
)
from tiktok_trip.export.map_data import build_map_data
from tiktok_trip.export.palette import (
    BLACK,
    CATEGORY_COLORS,
    CATEGORY_INK,
    DAY_COLORS,
    MIN_TEXT_CONTRAST,
    SECTOR_COLORS,
    SECTOR_INK,
    WHITE,
    contrast_ratio,
    readable_text_color,
    sector_ink,
)
from tiktok_trip.models import Hours, OpeningPeriod, UnplacedPlace
from tiktok_trip.planning.schema import UNPLACED_CODES


def every_day(open_: str, close: str | None) -> Hours:
    return Hours(periods=[OpeningPeriod(day=day, open=open_, close=close) for day in range(7)])


def test_contrast_ratio_matches_wcag() -> None:
    assert contrast_ratio(BLACK, WHITE) == pytest.approx(21.0)
    assert contrast_ratio("#777777", "#777777") == pytest.approx(1.0)
    assert contrast_ratio("#0072B2", WHITE) == pytest.approx(5.19, abs=0.01)


@pytest.mark.parametrize(
    ("background", "ink"),
    [("#009E73", BLACK), ("#D55E00", BLACK), ("#0072B2", WHITE), ("#6A3D9A", WHITE)],
)
def test_readable_text_color_uses_real_contrast(background: str, ink: str) -> None:
    assert readable_text_color(background) == ink


def test_every_sector_and_day_ink_reaches_aa() -> None:
    assert set(SECTOR_INK) >= set(SECTOR_COLORS)
    for sector, color in SECTOR_COLORS.items():
        assert contrast_ratio(SECTOR_INK[sector], color) >= MIN_TEXT_CONTRAST
    for color in DAY_COLORS:
        assert contrast_ratio(readable_text_color(color), color) >= MIN_TEXT_CONTRAST
    assert sector_ink(None) == SECTOR_INK["UNASSIGNED"]


def test_day_routes_stand_out_on_light_basemaps() -> None:
    for color in DAY_COLORS:
        assert contrast_ratio(color, WHITE) >= 3
        assert contrast_ratio(color, "#F2EFE9") >= 3
        assert color not in SECTOR_COLORS.values()


def test_hours_across_midnight() -> None:
    late = every_day("22:00", "02:00")
    assert opening_hours(EN, late) == f"10:00 PM{EN_DASH}2:00 AM daily"
    assert opening_hours(VI, late) == f"22:00{EN_DASH}2:00 hàng ngày"
    to_midnight = every_day("18:15", "00:00")
    assert opening_hours(EN, to_midnight) == f"6:15 PM{EN_DASH}12:00 AM daily"
    assert opening_hours(VI, to_midnight) == f"18:15{EN_DASH}24:00 hàng ngày"
    morning_to_night = every_day("10:00", "02:00")
    assert opening_hours(EN, morning_to_night) == f"10:00 AM{EN_DASH}2:00 AM daily"
    assert clock_range(EN, "08:12", "08:57") == f"8:12{EN_DASH}8:57 AM"


def test_open_24_hours_is_not_doubled() -> None:
    assert opening_hours(EN, every_day("00:00", None)) == "open 24 hours daily"
    assert opening_hours(VI, every_day("00:00", None)) == "mở cửa 24/24"
    assert opening_hours(EN, every_day("00:00", None), daily=False) == "open 24 hours"
    assert opening_hours(EN, every_day("00:00", "00:00")) == "open 24 hours daily"


def test_closed_days_are_named() -> None:
    hours = Hours(
        periods=[OpeningPeriod(day=day, open="11:30", close="18:30") for day in (0, 3, 4, 5, 6)]
    )
    assert opening_hours(EN, hours) == (
        f"Mon, Thu{EN_DASH}Sun 11:30 AM{EN_DASH}6:30 PM; closed Tue{EN_DASH}Wed"
    )
    assert opening_hours(VI, hours) == (
        f"T2, T5{EN_DASH}CN 11:30{EN_DASH}18:30; đóng cửa T3{EN_DASH}T4"
    )


def test_months_and_best_days_in_vietnamese() -> None:
    assert localize_calendar(VI, f"Nov{EN_DASH}Apr") == f"T11{EN_DASH}T4"
    assert localize_calendar(VI, "Sep, Oct") == "T9, T10"
    assert localize_calendar(VI, "weekend") == "cuối tuần"
    assert localize_calendar(EN, f"Nov{EN_DASH}Apr") == f"Nov{EN_DASH}Apr"


def test_nearby_kind_labels() -> None:
    assert kind_label(EN, "place_of_worship") == "Temple / pagoda"
    assert kind_label(VI, "place_of_worship") == "Chùa, đền"
    assert kind_label(EN, "city_gate") == "City gate"
    assert kind_label(VI, "city_gate") == VI.kinds["attraction"]


def test_plan_notes_in_both_languages() -> None:
    assert stop_reason(VI, "Bánh căn breakfast", "Ăn sáng bánh căn") == "Ăn sáng bánh căn"
    assert stop_reason(VI, "Bánh căn breakfast", None) == "Bánh căn breakfast"
    assert stop_reason(EN, "Bánh căn breakfast", "Ăn sáng bánh căn") == "Bánh căn breakfast"
    assert unplaced_text(EN, "too_far", "62 min away") == "Too far for this trip · 62 min away"
    assert unplaced_text(VI, "too_far", "62 min away") == "Quá xa cho chuyến này"
    assert unplaced_text(EN, None, "closed on Mondays") == "closed on Mondays"
    assert unplaced_text(EN, "too_far", "too far for this trip") == "Too far for this trip"
    for labels in (EN, VI):
        assert set(labels.unplaced_codes) == set(UNPLACED_CODES)


def test_stop_counts_have_singular_and_plural() -> None:
    assert EN.map_ui["stop_1"] == "{n} stop"
    assert EN.map_ui["stops"] == "{n} stops"
    assert VI.map_ui["stop_1"] == VI.map_ui["stops"] == "{n} điểm dừng"


def test_workbooks_use_vietnamese_notes_and_code_labels(tmp_path: Path) -> None:
    data = build_dataset()
    assert data.itinerary is not None
    first = data.itinerary.days[0].stops[1]
    first.reason_vi = "Phở buổi sáng"
    data.itinerary.unplaced = [UnplacedPlace(ref_id="gp:pho1", code="too_far", reason="45 km")]

    vi = load_workbook(write_trip(data, VI, tmp_path / "vi.xlsx"))
    en = load_workbook(write_trip(data, EN, tmp_path / "en.xlsx"))
    vi_rows = [[c.value for c in r] for r in vi[VI.plan_sheet].iter_rows(min_row=8)]
    en_rows = [[c.value for c in r] for r in en[EN.plan_sheet].iter_rows(min_row=8)]
    assert str(vi_rows[1][15]).startswith("Phở buổi sáng")
    assert str(en_rows[1][15]).startswith(str(first.reason))
    vi_out = next(i for i, r in enumerate(vi_rows) if r[0] == VI.not_planned)
    en_out = next(i for i, r in enumerate(en_rows) if r[0] == EN.not_planned)
    assert vi_rows[vi_out + 1][15] == "Quá xa cho chuyến này"
    assert en_rows[en_out + 1][15] == "Too far for this trip · 45 km"
    glance = [c.value for c in vi[VI.dashboard_sheet]["L"]]
    assert "Phở buổi sáng" in glance


def test_map_data_exposes_ink_kinds_notes_and_codes() -> None:
    data = build_dataset()
    assert data.itinerary is not None
    data.itinerary.days[0].stops[1].reason_vi = "Phở buổi sáng"
    data.itinerary.unplaced = [UnplacedPlace(ref_id="gp:pho1", code="closed")]

    map_data = build_map_data(data)

    assert map_data.sector_ink["E"] == BLACK
    assert map_data.sector_ink["N"] == WHITE
    assert len(map_data.day_ink) == len(map_data.day_colors)
    serialised = json.loads(map_data.model_dump_json(by_alias=True))
    assert serialised["category_colors"] == CATEGORY_COLORS
    assert serialised["category_ink"] == CATEGORY_INK
    museum = next(spot for spot in map_data.nearby if spot.kind == "museum")
    assert museum.kind_label == {"en": "Museum", "vi": "Bảo tàng"}
    assert map_data.days[0].stops[1].reason_vi == "Phở buổi sáng"
    unplaced = map_data.unplaced[0]
    assert unplaced.code == "closed"
    assert unplaced.reason_text == {
        "en": "Closed on the trip days",
        "vi": "Đóng cửa vào các ngày đi",
    }
    assert map_data.labels["vi"]["kinds"]["place_of_worship"] == "Chùa, đền"
    assert map_data.labels["en"]["unplaced_codes"]["low_priority"]
