"""Agent-mode route step: plan prompt for Claude Code, plan.json collect, clock and validation."""

from __future__ import annotations

import json
import os
from collections.abc import Sequence
from pathlib import Path

import pytest
from fixtures.sample_dataset import build_dataset
from openpyxl import load_workbook

from tiktok_trip import geo
from tiktok_trip.config import PlanConfig, TripConfig
from tiktok_trip.context import Context
from tiktok_trip.errors import GateError
from tiktok_trip.export.excel import write_trip
from tiktok_trip.export.i18n import EN, VI
from tiktok_trip.models import (
    BestTime,
    Category,
    Center,
    Hours,
    Itinerary,
    NearbySpot,
    OpeningPeriod,
    Place,
    Spatial,
    SpatialResult,
    UnplacedPlace,
    Zone,
)
from tiktok_trip.planning.prompt import target_per_day
from tiktok_trip.planning.schema import PlanParseError, parse_plan
from tiktok_trip.steps import route
from tiktok_trip.storage import load_model, save_model, save_models

CENTER = Center(lat=10.7769, lng=106.7009, source="hotel", label="Hotel")
KM_LAT = 1 / 111.195
MIN_PER_KM = 3


def _km(point: geo.LatLng) -> float:
    return round((point[0] - CENTER.lat) / KM_LAT, 3)


def fake_matrix(points: Sequence[geo.LatLng], mode: str = "motorbike") -> list[list[float]]:
    """3 minutes per km along the north-south line the test places sit on."""
    return [[abs(_km(a) - _km(b)) * MIN_PER_KM for b in points] for a in points]


def _place(
    pid: str, km_north: float, category: Category, hours: Hours | None = None, sector: str = "N"
) -> Place:
    return Place(
        id=pid,
        name=f"Place {pid}",
        category=category,
        address="x",
        lat=CENTER.lat + km_north * KM_LAT,
        lng=CENTER.lng,
        hours=hours,
        confidence=0.9,
        spatial=Spatial(
            bearing_deg=0 if km_north >= 0 else 180,
            sector=sector,
            distance_km=abs(km_north),
            ring="0-2 km",
            zone=f"{sector}-1",
        ),
    )


def _sat(open_: str, close: str) -> Hours:
    return Hours(periods=[OpeningPeriod(day=5, open=open_, close=close)])


PLACES = [
    _place("hill", 2, Category.ATTRACTION),
    _place("pho", 3, Category.FOOD, _sat("06:00", "11:00")),
    _place("lunch", 1, Category.FOOD),
    _place("cafe", 4, Category.CAFE),
    _place("dinner", -2, Category.FOOD, _sat("16:00", "22:00"), sector="S"),
    _place("far", 30, Category.ATTRACTION),
    _place(
        "monday",
        1.5,
        Category.OTHER,
        Hours(periods=[OpeningPeriod(day=0, open="08:00", close="17:00")]),
    ),
]
SPOT = NearbySpot(
    id="osm:node/7",
    name="Lake walk",
    kind="park",
    lat=CENTER.lat - 3 * KM_LAT,
    lng=CENTER.lng,
    zone_id="S-1",
    source="overpass",
)

GOOD_PLAN = {
    "days": [
        {
            "day": 1,
            "stops": [
                {"ref": "hill", "slot": "sunrise", "reason": "clouds at dawn"},
                {"ref": "pho", "slot": "breakfast", "reason": "phở after sunrise"},
                {"ref": "lunch", "slot": "lunch", "reason": "near the hotel"},
                {"ref": "rest", "slot": "rest", "reason": "nap"},
                {"ref": "cafe", "slot": "cafe", "reason": "view"},
            ],
        },
        {
            "day": 2,
            "stops": [
                {"ref": "dinner", "slot": "dinner", "reason": "lẩu"},
                {"ref": "osm:node/7", "slot": "evening", "reason": "walk after dinner"},
            ],
        },
    ],
    "unplaced": [
        {"ref": "far", "reason": "too far"},
        {"ref": "monday", "reason": "closed on Saturday"},
    ],
}


@pytest.fixture
def agent_ctx(ctx: Context, monkeypatch: pytest.MonkeyPatch) -> Context:
    north = [p.id for p in PLACES if p.id != "dinner"]
    zones = [
        Zone(id="N-1", sector="N", lat=CENTER.lat, lng=CENTER.lng, place_ids=north),
        Zone(id="S-1", sector="S", lat=CENTER.lat, lng=CENTER.lng, place_ids=["dinner"]),
    ]
    save_models(ctx.paths.places_file, PLACES)
    save_model(ctx.paths.spatial_file, SpatialResult(center=CENTER, sectors=8, zones=zones))
    save_models(ctx.paths.nearby_file, [SPOT])
    ctx.config.trip.days = 2
    monkeypatch.setattr(route, "travel_matrix", fake_matrix)
    return ctx


def write_plan_json(ctx: Context, plan: object, *, fenced: bool = False) -> None:
    text = json.dumps(plan, ensure_ascii=False)
    if fenced:
        text = f"Here is the plan:\n```json\n{text}\n```\n"
    ctx.paths.plan_json.write_text(text, encoding="utf-8")
    newer = ctx.paths.plan_prompt.stat().st_mtime + 5
    os.utime(ctx.paths.plan_json, (newer, newer))


def _gate(ctx: Context) -> GateError:
    with pytest.raises(GateError) as info:
        route.run(ctx)
    assert info.value.exit_code == 2
    return info.value


def test_default_plan_config() -> None:
    plan = PlanConfig()
    assert plan.mode == "agent"
    assert (plan.day_start, plan.day_end) == ("04:00", "22:00")
    assert [s.id for s in plan.rhythm] == [
        "sunrise",
        "breakfast",
        "morning",
        "lunch",
        "rest",
        "cafe",
        "dinner",
        "evening",
    ]
    assert plan.slot("rest") is not None
    assert plan.slot("rest").categories == []  # type: ignore[union-attr]
    assert plan.rest_at_center is True
    assert plan.target_places_per_day == 8
    morning = plan.slot("morning")
    assert morning is not None
    assert (morning.name_en, morning.name_vi) == ("Morning activity", "Hoạt động buổi sáng")


@pytest.mark.parametrize(
    ("total", "days", "cap", "expected"),
    [(176, 3, None, 8), (5, 3, None, 2), (0, 3, None, 0), (176, 3, 6, 6), (5, 1, 9, 5)],
)
def test_target_per_day(total: int, days: int, cap: int | None, expected: int) -> None:
    trip = TripConfig(max_places_per_day=cap)
    assert target_per_day(total, days, trip, PlanConfig()) == expected


def test_prompt_target_line_and_unplaced_sentence(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    prompt = agent_ctx.paths.plan_prompt.read_text(encoding="utf-8")
    per_day = target_per_day(
        len(PLACES), agent_ctx.config.trip.days, agent_ctx.config.trip, PlanConfig()
    )
    assert f"Aim for about {per_day} places a day" in prompt
    assert (
        "Places that do not fit go under `unplaced` with a code; prefer must, then high "
        "priority." in prompt
    )


def test_prompt_written_and_gate_until_plan_exists(agent_ctx: Context) -> None:
    error = _gate(agent_ctx)
    assert "waiting for Claude Code" in str(error)
    paths = agent_ctx.paths
    assert not paths.routes_file.exists()
    prompt = paths.plan_prompt.read_text(encoding="utf-8")
    for slot in PlanConfig().rhythm:
        assert f"`{slot.id}`" in prompt
        assert slot.name_en in prompt
    for place in PLACES:
        assert f"`{place.id}`" in prompt
    assert f"`{SPOT.id}`" in prompt
    assert str(paths.plan_matrix.resolve()) in prompt
    assert str(paths.plan_json.resolve()) in prompt
    assert '"unplaced"' in prompt and '"slot"' in prompt and '"reason"' in prompt
    assert "Do not compute or write any times" in prompt
    assert "04:00 to 22:00" in prompt
    assert "Hotel, the hotel" in prompt
    assert "CLOSED day 1, 2" in prompt
    assert "$" not in prompt.replace("$$", "")
    rows = paths.plan_matrix.read_text(encoding="utf-8").splitlines()
    header = rows[0].split(",")
    assert header[:2] == ["from", "center"]
    assert len(header) == 1 + 1 + len(PLACES) + 1
    assert len(rows) == len(header)
    assert rows[1].split(",")[header.index("hill")] == "6"


def test_prompt_lists_priority_first_and_states_the_rule(agent_ctx: Context) -> None:
    levels = {"far": "high", "dinner": "high", "cafe": "medium", "pho": "low"}
    places = [p.model_copy(update={"priority": levels.get(p.id)}) for p in PLACES]
    save_models(agent_ctx.paths.places_file, places)
    _gate(agent_ctx)
    prompt = agent_ctx.paths.plan_prompt.read_text(encoding="utf-8")

    assert "| id | name | priority | category |" in prompt
    assert "Then place every high-priority place; prefer medium over low when time" in prompt
    assert "medium and low places left out for time get the code `low_priority`" in prompt
    table = prompt[prompt.index("## Places") : prompt.index("## Nearby spots")]
    place_rows = [line for line in table.splitlines() if line.startswith("| `")]
    order = [r.split("`")[1] for r in place_rows]
    assert order[:4] == ["far", "dinner", "cafe", "pho"]
    assert set(order[4:]) == {"hill", "lunch", "monday"}
    assert place_rows[0].split(" | ")[2] == "high"
    assert place_rows[-1].split(" | ")[2] == "-"


def test_prompt_puts_must_places_first_and_leaves_removed_out(agent_ctx: Context) -> None:
    update = {"far": {"priority": "must"}, "monday": {"status": "skip", "preference": "removed"}}
    places = [p.model_copy(update=update.get(p.id, {})) for p in PLACES]
    save_models(agent_ctx.paths.places_file, places)
    _gate(agent_ctx)
    prompt = agent_ctx.paths.plan_prompt.read_text(encoding="utf-8")

    assert "You MUST place every `must` place" in prompt
    table = prompt[prompt.index("## Places") : prompt.index("## Nearby spots")]
    place_rows = [line for line in table.splitlines() if line.startswith("| `")]
    assert place_rows[0].split(" | ")[:3] == ["| `far`", "Place far", "must"]
    assert "`monday`" not in table


def test_unplaced_must_place_is_a_warning(agent_ctx: Context) -> None:
    places = [p.model_copy(update={"priority": "must"}) if p.id == "far" else p for p in PLACES]
    save_models(agent_ctx.paths.places_file, places)
    _gate(agent_ctx)
    write_plan_json(agent_ctx, GOOD_PLAN)

    route.run(agent_ctx)

    validation = agent_ctx.paths.plan_validation.read_text(encoding="utf-8")
    assert "must_unplaced" in validation
    assert "Place far (far) is a must-go place" in validation


def test_prompt_is_not_rewritten_when_nothing_changed(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    first = agent_ctx.paths.plan_prompt.stat().st_mtime_ns
    _gate(agent_ctx)
    assert agent_ctx.paths.plan_prompt.stat().st_mtime_ns == first


def test_collect_computes_the_clock(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    write_plan_json(agent_ctx, GOOD_PLAN, fenced=True)

    route.run(agent_ctx)

    itinerary = load_model(agent_ctx.paths.routes_file, Itinerary)
    assert itinerary.grouping == "agent"
    assert [(u.ref_id, u.reason) for u in itinerary.unplaced] == [
        ("far", "too far"),
        ("monday", "closed on Saturday"),
    ]
    day1, day2 = itinerary.days
    timeline = [(s.kind, s.ref_id, s.arrive, s.depart, s.travel_min_to_next) for s in day1.stops]
    assert timeline == [
        ("start", None, "04:24", "04:24", 6),
        ("place", "hill", "04:30", "05:45", 3),
        ("place", "pho", "07:00", "07:45", 6),
        ("place", "lunch", "11:30", "12:15", 3),
        ("rest", None, "12:18", "14:30", 12),
        ("place", "cafe", "14:42", "15:27", 12),
        ("end", None, "15:39", "15:39", None),
    ]
    assert day1.total_minutes == (15 * 60 + 39) - (4 * 60 + 24)
    assert [s.slot for s in day1.stops] == [
        None,
        "sunrise",
        "breakfast",
        "lunch",
        "rest",
        "cafe",
        None,
    ]
    assert day1.stops[1].reason == "clouds at dawn"
    assert day1.stops[4].name == "Hotel"
    assert day1.zone_ids == ["N-1"]
    assert day1.maps_urls
    timeline = [(s.kind, s.arrive, s.depart) for s in day2.stops]
    assert timeline == [
        ("start", "17:24", "17:24"),
        ("place", "17:30", "18:15"),
        ("nearby", "19:30", "20:15"),
        ("end", "20:24", "20:24"),
    ]
    validation = agent_ctx.paths.plan_validation.read_text(encoding="utf-8")
    assert "Status: **OK**" in validation
    assert "[long_wait]" in validation


def test_parse_plan_reads_vietnamese_notes_and_tolerates_codes() -> None:
    plan = parse_plan(
        json.dumps(
            {
                "days": [{"stops": [{"ref": "pho", "reason": "Phở", "reason_vi": "Ăn phở"}]}],
                "unplaced": [
                    {"ref": "a", "code": "too_far", "reason": "45 km"},
                    {"ref": "b", "code": "Out of days"},
                    {"ref": "c", "code": "sold out"},
                    {"ref": "d", "reason": "closed Mondays"},
                    {"ref": "e", "code": None, "reason": None},
                ],
            }
        )
    )
    assert plan.days[0].stops[0].reason_vi == "Ăn phở"
    assert [(u.code, u.reason) for u in plan.unplaced] == [
        ("too_far", "45 km"),
        ("out_of_days", ""),
        ("other", ""),
        (None, "closed Mondays"),
        (None, ""),
    ]


def test_collect_carries_vietnamese_notes_and_unplaced_codes(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    plan = json.loads(json.dumps(GOOD_PLAN))
    plan["days"][0]["stops"][0]["reason_vi"] = "Săn mây lúc bình minh"
    plan["unplaced"] = [
        {"ref": "far", "code": "too_far", "reason": "45 km"},
        {"ref": "monday", "reason": "closed on Saturday"},
    ]
    write_plan_json(agent_ctx, plan)

    route.run(agent_ctx)

    itinerary = load_model(agent_ctx.paths.routes_file, Itinerary)
    hill = itinerary.days[0].stops[1]
    assert (hill.reason, hill.reason_vi) == ("clouds at dawn", "Săn mây lúc bình minh")
    assert itinerary.days[0].stops[2].reason_vi is None
    assert [(u.ref_id, u.code, u.reason) for u in itinerary.unplaced] == [
        ("far", "too_far", "45 km"),
        ("monday", None, "closed on Saturday"),
    ]
    validation = agent_ctx.paths.plan_validation.read_text(encoding="utf-8")
    assert "Status: **OK**" in validation
    assert validation.count("[unplaced_code]") == 1
    assert "Place monday (monday) is under `unplaced` without a code" in validation


def test_prompt_asks_for_traveller_notes_in_both_languages(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    prompt = agent_ctx.paths.plan_prompt.read_text(encoding="utf-8")
    assert '"reason_vi"' in prompt and '"code"' in prompt
    assert "out_of_days | too_far | low_priority | closed | out_of_window | other" in prompt
    assert 'no "MUST-GO:"' in prompt
    assert "At most one breakfast stop a day" in prompt
    assert "never a second `dinner`" in prompt
    assert "back at the center by 21:00" in prompt


def test_violations_keep_old_routes(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    write_plan_json(agent_ctx, GOOD_PLAN)
    route.run(agent_ctx)
    before = agent_ctx.paths.routes_file.read_bytes()

    agent_ctx.config.plan.day_end = "20:00"
    _gate(agent_ctx)
    bad = json.loads(json.dumps(GOOD_PLAN))
    bad["days"][0]["stops"].append({"ref": "monday", "slot": "evening", "reason": "x"})
    bad["days"][0]["stops"].append({"ref": "gp:nope", "slot": "evening", "reason": "x"})
    bad["days"][1]["stops"].insert(0, {"ref": "pho", "slot": "breakfast", "reason": "again"})
    bad["unplaced"] = [{"ref": "far", "reason": "too far"}]
    write_plan_json(agent_ctx, bad)

    error = _gate(agent_ctx)

    assert agent_ctx.paths.routes_file.read_bytes() == before
    validation = agent_ctx.paths.plan_validation.read_text(encoding="utf-8")
    assert "Status: **FAILED**" in validation
    for rule in ("closed", "day_end", "duplicate", "unknown_ref"):
        assert f"[{rule}]" in validation, rule
        assert f"[{rule}]" in str(error)
    assert "Fix: " in validation
    assert "Place pho (pho) is planned 2 times" in validation


def test_missing_place_and_too_many_days(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    plan = json.loads(json.dumps(GOOD_PLAN))
    plan["unplaced"] = []
    plan["days"].append({"day": 3, "stops": [{"ref": "far", "slot": "morning"}]})
    write_plan_json(agent_ctx, plan)
    error = str(_gate(agent_ctx))
    assert "[missing_place]" in error
    assert "[too_many_days]" in error
    assert not agent_ctx.paths.routes_file.exists()


def test_stale_plan_waits_and_collect_flag_forces(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    write_plan_json(agent_ctx, GOOD_PLAN)
    route.run(agent_ctx)
    old = agent_ctx.paths.plan_prompt.stat().st_mtime - 10
    os.utime(agent_ctx.paths.plan_json, (old, old))
    assert "is older than the prompt" in str(_gate(agent_ctx))
    route.run(agent_ctx, collect=True)


def test_changed_places_regenerate_prompt_and_wait(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    write_plan_json(agent_ctx, GOOD_PLAN)
    route.run(agent_ctx)
    save_models(agent_ctx.paths.places_file, [*PLACES, _place("new", 0.5, Category.CAFE)])
    _gate(agent_ctx)
    assert "`new`" in agent_ctx.paths.plan_prompt.read_text(encoding="utf-8")


def test_unreadable_plan_is_reported(agent_ctx: Context) -> None:
    _gate(agent_ctx)
    write_plan_json(agent_ctx, {"days": [{"stops": [{"slot": "lunch"}]}]})
    assert "could not be read" in str(_gate(agent_ctx))
    assert "FAILED" in agent_ctx.paths.plan_validation.read_text(encoding="utf-8")
    with pytest.raises(PlanParseError):
        parse_plan("no json here")
    assert len(parse_plan('[{"stops": []}]').days) == 1


def test_solver_mode_is_unchanged(agent_ctx: Context) -> None:
    agent_ctx.config.plan.mode = "solver"
    route.run(agent_ctx)
    itinerary = load_model(agent_ctx.paths.routes_file, Itinerary)
    assert itinerary.grouping == "sector"
    assert not agent_ctx.paths.plan_prompt.exists()
    stops = [s for d in itinerary.days for s in d.stops]
    assert any(s.kind == "place" for s in stops)
    assert all(s.slot is None and s.kind != "rest" for s in stops)


def test_export_activity_column_and_rest_row(tmp_path: Path) -> None:
    data = build_dataset()
    assert data.itinerary is not None
    data.itinerary.unplaced = [UnplacedPlace(ref_id="gp:pho1", reason="closed")]
    first = data.itinerary.days[0].stops[1]
    assert first.reason is not None

    en = load_workbook(write_trip(data, EN, tmp_path / "en.xlsx"))[EN.plan_sheet]
    vi = load_workbook(write_trip(data, VI, tmp_path / "vi.xlsx"))[VI.plan_sheet]

    assert en["C6"].value == "Activity"
    assert vi["C6"].value == "Hoạt động"
    assert "planned by Claude Code" in en["A1"].value
    assert "lập bởi Claude Code" in vi["A1"].value
    en_rows = [[c.value for c in r] for r in en.iter_rows(min_row=8)]
    vi_rows = [[c.value for c in r] for r in vi.iter_rows(min_row=8)]
    assert en_rows[1][2] == "Breakfast"
    assert vi_rows[1][2] == "Ăn sáng"
    assert str(en_rows[1][15]).startswith(first.reason)
    assert en_rows[4][1] is None
    assert en_rows[4][5] == "Rest"
    assert en_rows[4][2] == "Relax at the hotel"
    assert str(en_rows[4][3]).endswith(',"Rest at center")')
    assert en_rows[4][7:10] == ["1:00 PM", "2:30 PM", 90]
    assert str(vi_rows[4][3]).endswith(',"Nghỉ ngơi tại trung tâm")')
    assert vi_rows[4][5] == "Nghỉ ngơi"
    assert en_rows[0][2] is None
    not_planned = next(i for i, r in enumerate(en_rows) if r[0] == "Not planned")
    assert str(en_rows[not_planned + 1][3]).endswith(',"Phở Hòa Pasteur")')
    assert en_rows[not_planned + 1][15] == "closed"
    assert data.spatial is not None
    data.spatial.center.source = "hotel"
    hotel = load_workbook(write_trip(data, VI, tmp_path / "hotel.xlsx"))[VI.plan_sheet]
    assert str(hotel["D12"].value).endswith(',"Nghỉ ngơi tại khách sạn")')


def test_prompt_shows_tags_and_best_time_reason(agent_ctx: Context) -> None:
    best = BestTime(hours="06:30-09:00", source="heuristic", why="breakfast place, early morning")
    update = {
        "pho": {"tags": ["breakfast"], "best_time": best},
        "cafe": {"tags": ["view", "sunset"]},
    }
    places = [p.model_copy(update=update.get(p.id, {})) for p in PLACES]
    save_models(agent_ctx.paths.places_file, places)
    _gate(agent_ctx)
    prompt = agent_ctx.paths.plan_prompt.read_text(encoding="utf-8")

    header = next(line for line in prompt.splitlines() if line.startswith("| id | name |"))
    assert header.endswith(
        "| best (source: why) | deadline | price | must try | note | mentions | tags |"
    )
    rows = {line.split("`")[1]: line for line in prompt.splitlines() if line.startswith("| `")}
    assert "| 06:30-09:00 (heuristic: breakfast place, early morning) |" in rows["pho"]
    assert rows["pho"].endswith("| breakfast |")
    assert rows["cafe"].endswith("| view, sunset |")
    assert rows["lunch"].endswith("| - |")
    assert "Match meal tags to meal slots" in prompt
    assert "cafés with view/sunset to the afternoon café slot" in prompt
    assert "sunrise/cloud_hunting to the sunrise slot" in prompt
