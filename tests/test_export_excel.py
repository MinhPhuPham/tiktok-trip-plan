# ruff: noqa: RUF001
"""trip.xlsx, one per language: the Dashboard sheet (title, totals, compass, plan at a glance)."""

from __future__ import annotations

from pathlib import Path

import pytest
from fixtures.sample_dataset import (
    CENTER_LABEL,
    CENTER_LABEL_EN,
    CENTER_PLACE_ID,
    LONG_REASON,
    build_dataset,
)
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from tiktok_trip.export.dataset import ExportData, first_sentence, truncate
from tiktok_trip.export.excel import trip_sheets, write_trip
from tiktok_trip.export.excel_style import set_link
from tiktok_trip.export.i18n import (
    EN,
    VI,
    Labels,
    clock,
    hours_and_best,
    localize_times,
    opening_hours,
)
from tiktok_trip.export.palette import CENTER_COLOR, SECTOR_COLORS, day_color
from tiktok_trip.models import BestTime, Hours, OpeningPeriod

SPACER_COL = 8
GLANCE = "IJKL"


@pytest.fixture
def data() -> ExportData:
    return build_dataset()


def trip_for(data: ExportData, labels: Labels, tmp_path: Path) -> Workbook:
    return load_workbook(write_trip(data, labels, tmp_path / labels.code / "trip.xlsx"))


def dashboard_for(data: ExportData, labels: Labels, tmp_path: Path) -> Worksheet:
    return trip_for(data, labels, tmp_path).worksheets[0]


@pytest.fixture
def workbook(data: ExportData, tmp_path: Path) -> Workbook:
    return trip_for(data, EN, tmp_path)


@pytest.fixture
def vi_workbook(data: ExportData, tmp_path: Path) -> Workbook:
    return trip_for(data, VI, tmp_path)


def test_two_sheets_per_language(workbook: Workbook, vi_workbook: Workbook) -> None:
    assert workbook.sheetnames == list(trip_sheets(EN)) == ["Dashboard", "Plan-to-go"]
    assert vi_workbook.sheetnames == list(trip_sheets(VI)) == ["Tổng Hợp", "Lịch Trình"]
    assert workbook.active.title == "Dashboard"
    assert vi_workbook.active.title == "Tổng Hợp"


SECTOR_HEADERS = {
    "B4": "NW (0)",
    "D4": "N (3)",
    "F4": "NE (2)",
    "B8": "W (2)",
    "F8": "E (1)",
    "B12": "SW (1)",
    "D12": "S (2)",
    "F12": "SE (1)",
}


def test_compass_blocks_are_two_columns_wide(workbook: Workbook) -> None:
    ws = workbook["Dashboard"]
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    for coordinate, text in SECTOR_HEADERS.items():
        assert ws[coordinate].value == text
        sector = text.split()[0]
        assert ws[coordinate].fill.fgColor.rgb.endswith(SECTOR_COLORS[sector].lstrip("#"))
        cell = ws[coordinate]
        assert f"{coordinate}:{ws.cell(row=cell.row, column=cell.column + 1).coordinate}" in merged
    assert ws["F12"].font.color.rgb.endswith("000000")
    assert ws["D4"].font.color.rgb.endswith("FFFFFF")
    assert ws["B5"].value is None


def test_compass_place_rows_name_and_combined_times(workbook: Workbook) -> None:
    ws = workbook["Dashboard"]
    assert str(ws["F13"].value).startswith('=HYPERLINK("https://www.google.com/maps/search/')
    assert ws["F13"].value.endswith(',"Bún Bò Huế O Hương")')
    assert ws["F13"].font.u == "single"
    assert ws["G13"].value == "7:00 AM–9:00 PM (best: 6:30–9:00 AM, estimate)"
    assert ws["D5"].value.endswith(',"Phở Hòa Pasteur")')
    assert ws["D6"].value.endswith(',"Bánh Mì Huỳnh Hoa")')
    assert ws["E6"].value == (
        "Mon–Fri 6:00 AM–2:00 PM; Sat–Sun 6:30 AM–12:00 PM (best: 5:00–8:00 PM)"
    )
    assert ws["D7"].value.endswith(',"Cà Phê Muối")')
    assert ws["E7"].value == "Tue–Sun 10:00 AM–2:00 PM, 5:00–10:00 PM; closed Mon"
    assert ws["G6"].value == "open 24 hours (best: 10:00 AM–12:00 PM)"
    assert ws["E14"].value == "7:00 AM–9:00 PM (best: 5:00–8:00 PM)"
    assert ws["C13"].value == "7:00 AM–9:00 PM"
    assert ws["E6"].alignment.wrap_text is True


def test_compass_vietnamese_uses_24_hour_clock(vi_workbook: Workbook) -> None:
    ws = vi_workbook["Tổng Hợp"]
    assert ws["F12"].value == "SE (1)"
    assert ws["G13"].value == "7:00–21:00 (giờ đẹp: 6:30–9:00, ước tính)"
    assert ws["E6"].value == "T2–T6 6:00–14:00; T7–CN 6:30–12:00 (giờ đẹp: 17:00–20:00)"
    assert ws["E14"].value == "7:00–21:00 (giờ đẹp: 17:00–20:00)"
    assert ws["G6"].value == "mở cửa 24 giờ (giờ đẹp: 10:00–12:00)"


def test_hours_and_best_combinations() -> None:
    daily = Hours(periods=[OpeningPeriod(day=d, open="06:30", close="15:00") for d in range(7)])
    best = BestTime(hours="07:00-09:00", source="comments")
    assert hours_and_best(EN, daily, best) == "6:30 AM–3:00 PM (best: 7:00–9:00 AM)"
    assert hours_and_best(VI, daily, best) == "6:30–15:00 (giờ đẹp: 7:00–9:00)"
    assert hours_and_best(EN, daily, None) == "6:30 AM–3:00 PM"
    assert hours_and_best(EN, daily, BestTime(days="weekend", source="caption")) == (
        "6:30 AM–3:00 PM"
    )
    assert hours_and_best(EN, None, best) == "Best time: 7:00–9:00 AM"
    assert hours_and_best(VI, None, best) == "Giờ đẹp: 7:00–9:00"
    assert hours_and_best(EN, None, None) == "—"
    assert hours_and_best(VI, Hours(), None) == "—"


def test_center_block_links_to_maps(workbook: Workbook, vi_workbook: Workbook) -> None:
    ws = workbook["Dashboard"]
    assert ws["D8"].value == "CENTER"
    assert ws["D8"].fill.fgColor.rgb.endswith(CENTER_COLOR.lstrip("#"))
    center = ws["D9"].value
    assert center.startswith('=HYPERLINK("https://www.google.com/maps/search/')
    assert f"query_place_id={CENTER_PLACE_ID}" in center
    assert center.endswith(f',"{CENTER_LABEL_EN}")')
    assert ws["D10"].value == "Source: City default"
    assert ws["D11"].value == "10.77210, 106.69800"
    assert {"D9:E9", "D10:E10", "D11:E11"} <= {str(rng) for rng in ws.merged_cells.ranges}
    vi = vi_workbook["Tổng Hợp"]
    assert vi["D8"].value == "TRUNG TÂM"
    assert vi["D9"].value.endswith(f',"{CENTER_LABEL}")')
    assert vi["D10"].value == "Nguồn: Mặc định của thành phố"


def test_center_without_english_label_falls_back(data: ExportData, tmp_path: Path) -> None:
    assert data.spatial is not None
    data.spatial.center = data.spatial.center.model_copy(
        update={"label_en": None, "place_id": None}
    )
    ws = dashboard_for(data, EN, tmp_path)
    assert ws["D9"].value.endswith(f',"{CENTER_LABEL}")')
    assert "query=10.7721,106.698" in ws["D9"].value


def glance_row(ws, row: int) -> list:  # type: ignore[no-untyped-def]
    return [ws[f"{col}{row}"].value for col in GLANCE]


def test_glance_header_and_day_blocks(workbook: Workbook) -> None:
    ws = workbook["Dashboard"]
    assert glance_row(ws, 4) == ["Activity", "Place", "Time", "Note"]
    assert ws["I4"].fill.fgColor.rgb.endswith("37474F")
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert {"I5:L5", "I14:L14"} <= merged
    assert ws["I5"].value == "Day 1 · N, NE · N-1, NE-1 · 8:00 AM–8:05 PM"
    assert ws["I14"].value == "Day 2 · S, SW · S-1, SW-1 · 8:00 AM–8:55 PM"
    for cell, day in (("I5", 1), ("I14", 2)):
        assert ws[cell].fill.fgColor.rgb.endswith(day_color(day).lstrip("#"))
        assert ws[cell].font.b is True
    assert [ws.column_dimensions[letter].width for letter in GLANCE] == [24, 32, 17, 42]
    assert ws.column_dimensions["H"].width == 3
    for row in range(4, 21):
        assert ws.cell(row=row, column=SPACER_COL).value is None
    assert ws.auto_filter.ref is None
    assert ws.print_area == "'Dashboard'!$A$1:$L$20"


def test_glance_stop_rows(workbook: Workbook) -> None:
    ws = workbook["Dashboard"]
    activities = [ws[f"I{row}"].value for row in range(6, 14)]
    assert activities == [
        "Start",
        "Breakfast",
        "Morning activity",
        "Sight",
        "Relax at the hotel",
        "Dinner",
        "Hang out",
        "End",
    ]
    assert [ws[f"K{row}"].value for row in range(6, 14)] == [
        "8:00 AM",
        "8:10–8:55 AM",
        "9:03–9:30 AM",
        "9:42–10:15 AM",
        "1:00–2:30 PM",
        "5:30–6:15 PM",
        "6:35–7:50 PM",
        "8:05 PM",
    ]
    assert ws["J7"].value.startswith('=HYPERLINK("https://www.google.com/maps/search/')
    assert ws["J7"].value.endswith(',"Phở Hòa Pasteur")')
    assert ws["J7"].font.u == "single"
    assert ws["J9"].value.endswith(',"Nhà Thờ Đức Bà")')
    for row in (6, 10, 13):
        center = ws[f"J{row}"].value
        assert f"query_place_id={CENTER_PLACE_ID}" in center
        assert center.endswith(f',"{CENTER_LABEL_EN}")')
    assert ws["L8"].value == "Bánh mì to go"
    assert ws["L6"].value is None
    assert ws["L10"].value is None
    assert ws["L7"].alignment.wrap_text is True


def test_glance_note_is_reason_cut_at_60(workbook: Workbook) -> None:
    note = workbook["Dashboard"]["L7"].value
    assert len(LONG_REASON) > 60
    assert note == "Breakfast phở before the queue; locals say the broth sells…"
    assert len(note) <= 60
    assert truncate(LONG_REASON, 60) == note
    assert truncate("short", 60) == "short"


def test_glance_vietnamese(vi_workbook: Workbook) -> None:
    ws = vi_workbook["Tổng Hợp"]
    assert glance_row(ws, 4) == ["Hoạt động", "Địa điểm", "Thời gian", "Ghi chú"]
    assert ws["I5"].value == "Ngày 1 · N, NE · N-1, NE-1 · 8:00–20:05"
    assert glance_row(ws, 6)[0] == "Xuất phát"
    assert glance_row(ws, 6)[2] == "8:00"
    assert ws["J6"].value.endswith(f',"{CENTER_LABEL}")')
    assert glance_row(ws, 7)[0] == "Ăn sáng"
    assert glance_row(ws, 7)[2] == "8:10–8:55"
    assert glance_row(ws, 10)[0] == "Nghỉ ngơi ở khách sạn"
    assert glance_row(ws, 10)[2] == "13:00–14:30"
    assert glance_row(ws, 13)[0] == "Kết thúc"
    assert ws["I14"].value == "Ngày 2 · S, SW · S-1, SW-1 · 8:00–20:55"
    assert ws.print_area == "'Tổng Hợp'!$A$1:$L$20"


def test_glance_without_slots_falls_back_to_category(data: ExportData, tmp_path: Path) -> None:
    assert data.itinerary is not None
    for day in data.itinerary.days:
        for stop in day.stops:
            stop.slot = None
    ws = dashboard_for(data, EN, tmp_path)
    assert [ws[f"I{row}"].value for row in range(6, 14)] == [
        "Start",
        "Food",
        "Food",
        "Sight",
        "Rest",
        "Food",
        "Attraction",
        "End",
    ]


def test_glance_without_itinerary_shows_one_line(data: ExportData, tmp_path: Path) -> None:
    data.itinerary = None
    ws = dashboard_for(data, EN, tmp_path)
    assert ws["I4"].value == "Plan not generated yet: run tiktok-trip route"
    assert all(ws.cell(row=row, column=9).value is None for row in range(5, ws.max_row + 1))
    assert not [rng for rng in ws.merged_cells.ranges if rng.min_col >= 9]
    vi = dashboard_for(data, VI, tmp_path)
    assert vi["I4"].value == VI.no_plan == "Chưa có lịch trình: hãy chạy tiktok-trip route"


def test_row_heights_capped(data: ExportData, tmp_path: Path) -> None:
    assert data.itinerary is not None
    data.itinerary.days[0].stops[2].reason = "word " * 40
    ws = dashboard_for(data, EN, tmp_path)
    heights = [ws.row_dimensions[row].height or 0 for row in range(1, ws.max_row + 1)]
    assert max(heights) <= 45
    assert ws.row_dimensions[6].height == 45
    assert ws.row_dimensions[8].height == 30
    assert len(ws["L8"].value) <= 60


def test_dashboard_title_totals_and_print(workbook: Workbook, vi_workbook: Workbook) -> None:
    ws = workbook["Dashboard"]
    assert ws["B1"].value == "Ho Chi Minh City trip plan — generated 2026-10-19"
    assert ws["B2"].value == (
        "Places 12 · Zones 8 · Videos 16 · Pending review 2 · High priority 0"
    )
    assert "B2:G2" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["B2"].alignment.horizontal == "left"
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    vi = vi_workbook["Tổng Hợp"]
    assert vi["B1"].value == "Kế hoạch chuyến đi Ho Chi Minh City — tạo ngày 19/10/2026"
    assert vi["B2"].value == "Địa điểm 12 · Khu 8 · Video 16 · Chờ duyệt 2 · Ưu tiên cao 0"


def test_places_without_direction(data: ExportData, tmp_path: Path) -> None:
    data.places[0] = data.places[0].model_copy(update={"spatial": None})
    dashboard = dashboard_for(data, EN, tmp_path)
    assert dashboard["B2"].value == (
        "Places 12 · Without direction 1 · Zones 8 · Videos 16 · Pending review 2 · High priority 0"
    )
    assert dashboard["D4"].value == "N (2)"


def test_first_sentence_cut() -> None:
    long_note = (
        "Mở 6:30-15:00 (info card 069.jpg; caption 'mở tới 15h chiều'); quán khuất, view rừng "
        "thông (transcript 00:00-00:04); xe hơi vào tận quán. Second sentence is dropped."
    )
    body = first_sentence([long_note, "Other note"])
    assert len(body) <= 90
    assert body.endswith("rừng…")
    assert first_sentence(["Short one. Then more."]) == "Short one."
    assert first_sentence([]) == ""


def test_long_url_falls_back_to_plain_hyperlink(data: ExportData, tmp_path: Path) -> None:
    assert data.itinerary is not None
    long_name = "Phở " + "x" * 300
    data.places[0] = data.places[0].model_copy(update={"name": long_name})
    ws = dashboard_for(data, EN, tmp_path)
    assert ws["J7"].value == long_name
    assert ws["J7"].hyperlink.target.startswith("https://www.google.com/maps/search/")


def test_set_link_long_url_is_plain_clickable_hyperlink() -> None:
    url = "https://www.google.com/maps/search/?api=1&query=" + "a" * 260
    assert len(url) >= 300
    cell = Workbook().active["A1"]
    set_link(cell, url, "★ Đường Sương Nguyệt Anh")
    assert cell.value == "★ Đường Sương Nguyệt Anh"
    assert cell.hyperlink.target == url
    assert cell.font.underline == "single"
    assert cell.font.color.rgb.endswith("0563C1")


def test_compass_and_plan_rows_keep_link_for_long_url(data: ExportData, tmp_path: Path) -> None:
    place = data.places[0].model_copy(update={"id": "gp:" + "E" * 300})
    data.places[0] = place
    assert len(place.maps_url) > 300
    wb = trip_for(data, EN, tmp_path)
    linked = [
        cell
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.hyperlink is not None and cell.hyperlink.target == place.maps_url
    ]
    assert linked
    assert all(isinstance(cell.value, str) and place.name in cell.value for cell in linked)


def test_four_sector_dashboard_uses_cardinals_only(tmp_path: Path) -> None:
    ws = dashboard_for(build_dataset(sectors=4), EN, tmp_path)
    headers = {str(rng) for rng in ws.merged_cells.ranges if rng.min_col == rng.max_col - 1}
    assert {"D4:E4", "B8:C8", "D8:E8", "F8:G8", "D12:E12"} <= headers
    assert ws["B4"].value is None
    assert ws["F4"].value is None
    assert ws["D4"].value.startswith("N (")
    assert ws["B8"].value.startswith("W (")
    assert ws["F8"].value.startswith("E (")
    assert ws["D12"].value.startswith("S (")


def test_block_height_grows_with_busiest_sector(data: ExportData, tmp_path: Path) -> None:
    extra = [
        place.model_copy(update={"id": f"{place.id}-copy{n}", "name": f"{place.name} {n}"})
        for n, place in enumerate(data.places[:2])
    ]
    data.places.extend(extra)
    ws = dashboard_for(data, EN, tmp_path)
    assert ws["D4"].value == "N (5)"
    assert ws["B10"].value == "W (2)"
    assert ws["D10"].value == "CENTER"
    assert ws["F16"].value == "SE (1)"


def test_clock_styles() -> None:
    assert clock(EN, "07:30") == "7:30 AM"
    assert clock(EN, "00:00") == "12:00 AM"
    assert clock(EN, "12:15") == "12:15 PM"
    assert clock(EN, "24:00") == "12:00 AM"
    assert clock(VI, "07:30") == "7:30"
    assert clock(VI, "15:00") == "15:00"
    assert localize_times(EN, "06:30–15:00") == "6:30 AM–3:00 PM"
    assert (
        localize_times(VI, "around 08:00, before 10:00-11:30") == "around 8:00, before 10:00–11:30"
    )
    assert localize_times(EN, "no times here") == "no times here"


def test_opening_hours_short_forms() -> None:
    daily = Hours(periods=[OpeningPeriod(day=d, open="06:30", close="15:00") for d in range(7)])
    assert opening_hours(EN, daily) == "6:30 AM–3:00 PM daily"
    assert opening_hours(VI, daily) == "6:30–15:00 hàng ngày"
    gaps = Hours(
        periods=[OpeningPeriod(day=d, open="08:00", close="17:00") for d in (0, 1, 2, 4, 5)]
    )
    assert opening_hours(EN, gaps) == "Mon–Wed, Fri–Sat 8:00 AM–5:00 PM; closed Thu, Sun"
    late = Hours(periods=[OpeningPeriod(day=6, open="18:00", close=None)])
    assert opening_hours(VI, late) == "CN từ 18:00; đóng cửa T2–T7"
    text_only = Hours(weekday_text=["Monday: 07:00 – 21:00"])
    assert opening_hours(EN, text_only) == "Monday: 7:00 AM–9:00 PM"
    assert opening_hours(EN, None) == ""


def test_opening_hours_without_daily_word() -> None:
    daily = Hours(periods=[OpeningPeriod(day=d, open="06:30", close="15:00") for d in range(7)])
    assert opening_hours(EN, daily, daily=False) == "6:30 AM–3:00 PM"
    assert opening_hours(VI, daily, daily=False) == "6:30–15:00"
    weekdays = Hours(periods=[OpeningPeriod(day=d, open="08:00", close="17:00") for d in range(5)])
    assert opening_hours(EN, weekdays, daily=False) == "Mon–Fri 8:00 AM–5:00 PM; closed Sat–Sun"


def _with_priorities(data: ExportData, levels: dict[str, str], videos: dict[str, int]) -> None:
    for index, place in enumerate(data.places):
        slug = place.id.removeprefix("gp:")
        update: dict[str, object] = {"priority": levels.get(slug)}
        if slug in videos:
            update["videos"] = [
                f"https://www.tiktok.com/@r/video/{slug}{n}" for n in range(videos[slug])
            ]
        data.places[index] = place.model_copy(update=update)


def test_compass_orders_by_priority_and_stars_the_names(data: ExportData, tmp_path: Path) -> None:
    levels = {"cafe3": "high", "banhmi2": "medium", "che11": "low", "market12": "low"}
    _with_priorities(data, levels, {"market12": 2})
    ws = dashboard_for(data, EN, tmp_path)

    assert ws["D5"].value.endswith(',"★★★ Cà Phê Muối")')
    assert ws["E5"].value == "Tue–Sun 10:00 AM–2:00 PM, 5:00–10:00 PM; closed Mon"
    assert ws["D6"].value.endswith(',"★★ Bánh Mì Huỳnh Hoa")')
    assert ws["D7"].value.endswith(',"Phở Hòa Pasteur")')
    assert ws["B9"].value.endswith(',"★ Chợ Bình Tây")')
    assert ws["B10"].value.endswith(',"★ Chè Thái Lan")')
    assert ws["B2"].value.endswith("Pending review 2 · High priority 1")
    assert ws["B16"].value == "★★★ high · ★★ medium · ★ low"
    assert ws["B16"].font.i is True


def test_priority_legend_and_totals_in_vietnamese(data: ExportData, tmp_path: Path) -> None:
    _with_priorities(data, {"pho1": "high", "oc9": "high"}, {})
    ws = dashboard_for(data, VI, tmp_path)
    assert ws["B2"].value.endswith("Ưu tiên cao 2")
    assert ws["B16"].value == "★★★ cao · ★★ vừa · ★ thấp"
    assert ws["D13"].value.endswith(',"★★★ Ốc Đào")')
    assert ws["D14"].value.endswith(',"Hủ Tiếu Nam Vang")')
