# ruff: noqa: RUF001
"""trip.xlsx Plan-to-go sheet: day summary, every stop day by day, unplaced places."""

from __future__ import annotations

from pathlib import Path

import pytest
from fixtures.sample_dataset import (
    CENTER_LABEL,
    CENTER_LABEL_EN,
    CENTER_PLACE_ID,
    LONG_REASON,
    build_dataset,
)
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from tiktok_trip.export.dataset import ExportData
from tiktok_trip.export.excel import write_trip
from tiktok_trip.export.excel_plan import Lookups, areas_text
from tiktok_trip.export.i18n import EN, VI, Labels
from tiktok_trip.export.palette import day_color
from tiktok_trip.models import DayPlan

PLAN_HEADER = [
    "Day",
    "#",
    "Activity",
    "Place",
    "Priority",
    "Type",
    "Area",
    "Arrive",
    "Leave",
    "Time there (min)",
    "Travel to next (min)",
    "Open hours",
    "Best time",
    "Tags",
    "Price",
    "Note",
]
DETAIL_HEADER_ROW = 6


@pytest.fixture
def data() -> ExportData:
    return build_dataset()


def plan_for(data: ExportData, labels: Labels, tmp_path: Path) -> Worksheet:
    workbook = load_workbook(write_trip(data, labels, tmp_path / labels.code / "trip.xlsx"))
    return workbook[labels.plan_sheet]


@pytest.fixture
def ws(data: ExportData, tmp_path: Path) -> Worksheet:
    return plan_for(data, EN, tmp_path)


@pytest.fixture
def vi_ws(data: ExportData, tmp_path: Path) -> Worksheet:
    return plan_for(data, VI, tmp_path)


def values(ws: Worksheet, row: int) -> list[object]:
    return [cell.value for cell in ws[row]][:16]


def stop_rows(ws: Worksheet) -> list[list[object]]:
    return [
        values(ws, row)
        for row in range(DETAIL_HEADER_ROW + 1, ws.max_row + 1)
        if isinstance(ws.cell(row=row, column=1).value, int)
    ]


def test_title_and_summary(ws: Worksheet) -> None:
    assert ws["A1"].value == (
        "Ho Chi Minh City plan to go — 2 days · planned by Claude Code · motorbike"
    )
    header = {cell.coordinate: cell.value for cell in ws[2] if cell.value}
    assert header == {
        "A2": "Day",
        "B2": "Areas",
        "E2": "Stops",
        "F2": "Start",
        "G2": "End",
        "H2": "Total time (min)",
        "I2": "Travel time (min)",
        "J2": "Directions",
    }
    assert ws["K2"].fill.fgColor.rgb.endswith("37474F")
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert {"B2:D2", "B3:D3", "B4:D4"} <= merged
    day1, day2 = values(ws, 3), values(ws, 4)
    assert [day1[i] for i in (0, 1, 4, 5, 6, 7, 8)] == [
        "Day 1",
        "N, NE · N-1, NE-1",
        5,
        "8:00 AM",
        "8:05 PM",
        725,
        100,
    ]
    assert [day2[i] for i in (0, 1, 4, 5, 6, 7, 8)] == [
        "Day 2",
        "S, SW · S-1, SW-1",
        4,
        "8:00 AM",
        "8:55 PM",
        775,
        51,
    ]
    assert str(day1[9]).startswith('=HYPERLINK("https://www.google.com/maps/dir/')
    assert str(day1[9]).endswith(',"Directions")')
    assert day1[10] is None
    assert str(day2[9]).endswith(',"Leg 1")')
    assert str(day2[10]).endswith(',"Leg 2")')
    assert "destination=10.77,106.69" in str(day2[10])
    assert ws["A3"].fill.fgColor.rgb.endswith(day_color(1).lstrip("#"))
    assert ws["A4"].fill.fgColor.rgb.endswith(day_color(2).lstrip("#"))
    assert ws["H3"].number_format == "0"


def test_title_shows_solver_grouping(data: ExportData, tmp_path: Path) -> None:
    assert data.itinerary is not None
    data.itinerary.grouping = "nearest"
    assert (
        "grouped by nearest zones · optimised by solver" in plan_for(data, EN, tmp_path)["A1"].value
    )
    assert "chia theo khu gần nhau" in plan_for(data, VI, tmp_path)["A1"].value


def test_detail_header_frozen_and_filtered(ws: Worksheet) -> None:
    assert values(ws, DETAIL_HEADER_ROW) == PLAN_HEADER
    assert ws["A6"].fill.fgColor.rgb.endswith("37474F")
    assert ws.freeze_panes == "A7"
    assert ws.auto_filter.ref == "A6:P22"
    assert ws.print_title_rows == "$6:$6"
    assert ws["A7"].value == "Day 1 · N, NE · N-1, NE-1 · 8:00 AM–8:05 PM"
    assert ws["A16"].value == "Day 2 · S, SW · S-1, SW-1 · 8:00 AM–8:55 PM"
    for cell, day in (("A7", 1), ("P7", 1), ("A16", 2)):
        assert ws[cell].fill.fgColor.rgb.endswith(day_color(day).lstrip("#"))
    assert ws["A7"].font.b is True


def test_detail_stop_rows(ws: Worksheet) -> None:
    stops = stop_rows(ws)
    assert len(stops) == 8 + 6
    start, pho = stops[0], stops[1]
    assert start[:11] == [
        1,
        None,
        None,
        start[3],
        None,
        "Start",
        "Center",
        None,
        "8:00 AM",
        None,
        10,
    ]
    assert f"query_place_id={CENTER_PLACE_ID}" in str(start[3])
    assert str(start[3]).endswith(f',"{CENTER_LABEL_EN}")')
    assert pho[:3] == [1, 1, "Breakfast"]
    assert str(pho[3]).endswith(',"Phở Hòa Pasteur")')
    assert pho[4:15] == [
        None,
        "Place",
        "North of center · Zone N-1",
        "8:10 AM",
        "8:55 AM",
        45,
        8,
        "7:00 AM–9:00 PM daily",
        "6:30–9:00 AM (estimate)",
        None,
        "50k-100k VND",
    ]
    assert pho[15] == f"{LONG_REASON} · Comment about Phở Hòa Pasteur: queue after 18:00"
    nearby = stops[3]
    assert nearby[1:3] == [3, "Sight"]
    assert nearby[4:11] == [
        None,
        "Nearby spot",
        "North of center · Zone N-1",
        "9:42 AM",
        "10:15 AM",
        33,
        15,
    ]
    assert nearby[15] == "Two minutes from the bánh mì"
    end = stops[7]
    assert end[4:11] == [None, "End", "Center", "8:05 PM", None, None, None]
    assert [row[5] for row in stops[8:]] == [
        "Start",
        "Place",
        "Nearby spot",
        "Place",
        "Place",
        "End",
    ]
    assert [row[1] for row in stops[8:]] == [None, 1, 2, 3, 4, None]
    assert stops[12][6] == "Southwest of center · Zone SW-1"


def test_rest_row(ws: Worksheet, data: ExportData, tmp_path: Path) -> None:
    rest = stop_rows(ws)[4]
    assert rest[1] is None
    assert rest[2] == "Relax at the hotel"
    assert str(rest[3]).endswith(',"Rest at center")')
    assert rest[4:11] == [None, "Rest", "Center", "1:00 PM", "2:30 PM", 90, 20]
    assert rest[15] is None
    assert data.spatial is not None
    data.spatial.center.source = "hotel"
    vi_rest = stop_rows(plan_for(data, VI, tmp_path))[4]
    assert str(vi_rest[3]).endswith(',"Nghỉ ngơi tại khách sạn")')
    assert vi_rest[2] == "Nghỉ ngơi ở khách sạn"


def test_links_and_numeric_minutes(ws: Worksheet) -> None:
    for row in range(DETAIL_HEADER_ROW + 1, 23):
        if not isinstance(ws.cell(row=row, column=1).value, int):
            continue
        place = ws.cell(row=row, column=4)
        assert str(place.value).startswith('=HYPERLINK("https://www.google.com/maps/search/')
        assert place.font.u == "single"
        for col in (10, 11):
            cell = ws.cell(row=row, column=col)
            assert cell.value is None or isinstance(cell.value, int)
            assert cell.number_format == "0"


def test_unplaced_at_bottom(ws: Worksheet) -> None:
    assert ws["A24"].value == "Not planned"
    assert ws["A24"].font.b is True
    assert str(ws["D25"].value).endswith(',"Cơm Tấm Ba Ghiền")')
    assert ws["G25"].value == "East of center · Zone E-1"
    assert ws["P25"].value == "No lunch slot left near E-1"
    assert ws["D26"].value == "gp:gone"
    assert ws["P26"].value == "Closed for renovation"
    assert ws.max_row == 26


def test_vietnamese(vi_ws: Worksheet) -> None:
    assert vi_ws.title == "Lịch Trình"
    assert vi_ws["A1"].value == (
        "Lịch trình Ho Chi Minh City — 2 ngày · lập bởi Claude Code · xe máy"
    )
    assert values(vi_ws, 3)[:7] == ["Ngày 1", "N, NE · N-1, NE-1", None, None, 5, "8:00", "20:05"]
    assert str(vi_ws["K4"].value).endswith(',"Chặng 2")')
    assert values(vi_ws, DETAIL_HEADER_ROW) == list(VI.plan_columns)
    assert vi_ws["A7"].value == "Ngày 1 · N, NE · N-1, NE-1 · 8:00–20:05"
    stops = stop_rows(vi_ws)
    assert {row[5] for row in stops} == {
        "Xuất phát",
        "Địa điểm",
        "Điểm lân cận",
        "Nghỉ ngơi",
        "Kết thúc",
    }
    assert str(stops[0][3]).endswith(f',"{CENTER_LABEL}")')
    assert stops[0][6] == "Trung tâm"
    assert stops[1][2] == "Ăn sáng"
    assert stops[4][2] == "Nghỉ ngơi ở khách sạn"
    assert str(stops[4][3]).endswith(',"Nghỉ ngơi tại trung tâm")')
    assert stops[5][6:9] == ["Phía đông bắc trung tâm · Khu NE-1", "17:30", "18:15"]
    assert stops[1][11:13] == ["7:00–21:00 hàng ngày", "6:30–9:00 (ước tính)"]
    assert vi_ws["A24"].value == "Chưa xếp lịch"


def test_note_is_reason_then_first_sentence(data: ExportData, tmp_path: Path) -> None:
    long_note = "A" * 40 + " " + "b" * 70 + ". Second sentence."
    data.places[0] = data.places[0].model_copy(update={"notes": [long_note]})
    assert data.itinerary is not None
    data.itinerary.days[0].stops[1].reason = None
    note = str(stop_rows(plan_for(data, EN, tmp_path))[1][15])
    assert len(note) <= 90
    assert note.endswith("…")
    assert "Second" not in note


def test_without_itinerary(data: ExportData, tmp_path: Path) -> None:
    data.itinerary = None
    ws = plan_for(data, EN, tmp_path)
    assert ws["A3"].value == EN.no_plan
    assert "0 days" in ws["A1"].value
    assert values(ws, 5) == PLAN_HEADER
    assert ws.freeze_panes == "A6"
    assert ws.max_row == 5


def test_without_spatial_links_center_by_coordinates(tmp_path: Path) -> None:
    data = build_dataset()
    data.spatial = None
    stops = stop_rows(plan_for(data, EN, tmp_path))
    assert "query=10.7721,106.698" in str(stops[0][3])
    assert stops[0][2] is None
    assert str(stops[0][3]).endswith(f',"{CENTER_LABEL}")')


def test_areas_text_falls_back_to_zone_prefix(data: ExportData) -> None:
    lookups = Lookups.build(data)
    assert areas_text(DayPlan(day=1, zone_ids=["S-1", "SW-1", "S-2"]), lookups) == (
        "S, SW · S-1, SW-1, S-2"
    )
    assert areas_text(DayPlan(day=1), lookups) == ""
    assert EN.day_count(1) == "1 day"
    assert VI.day_count(3) == "3 ngày"


def test_priority_column_after_place(data: ExportData, tmp_path: Path) -> None:
    levels = {"gp:pho1": "high", "gp:banhmi2": "medium", "gp:comtam6": "low"}
    data.places = [p.model_copy(update={"priority": levels.get(p.id)}) for p in data.places]
    en, vi = plan_for(data, EN, tmp_path), plan_for(data, VI, tmp_path)

    assert en["E6"].value == "Priority"
    assert vi["E6"].value == "Ưu tiên"
    assert [row[4] for row in stop_rows(en)[:4]] == [None, "High", "Medium", None]
    assert [row[4] for row in stop_rows(vi)[:3]] == [None, "Cao", "Vừa"]
    assert en["E25"].value == "Low"
    assert vi["E25"].value == "Thấp"
    assert en["E26"].value is None
