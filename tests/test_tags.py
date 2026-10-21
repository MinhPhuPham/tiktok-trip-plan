# ruff: noqa: RUF001
"""Place tags (tags.py), the tag-driven best time with its reason, and how both are shown."""

from __future__ import annotations

from pathlib import Path

import pytest
from fixtures.sample_dataset import build_dataset
from openpyxl import load_workbook

from tiktok_trip.besttime import PlaceFacts, SourceText, suggest
from tiktok_trip.config import CityProfile
from tiktok_trip.export.excel import write_trip
from tiktok_trip.export.i18n import EN, VI, best_time_why, tags_text
from tiktok_trip.models import BestTime, Category, Hours, OpeningPeriod, Place
from tiktok_trip.steps import enrich
from tiktok_trip.tags import (
    BEST_TIME_ORDER,
    TAG_WHY,
    TAG_WINDOWS,
    TAGS,
    TagFacts,
    derive_tags,
    tag_hits,
)

DA_LAT = CityProfile(aliases=["Đà Lạt"], best_months="Nov–Apr", viewpoints=["Langbiang"])


def hours(open_: str, close: str | None) -> Hours:
    return Hours(periods=[OpeningPeriod(day=d, open=open_, close=close) for d in range(7)])


def facts(
    name: str,
    category: Category = Category.FOOD,
    *,
    must_try: tuple[str, ...] = (),
    notes: tuple[str, ...] = (),
    opening: Hours | None = None,
    best: tuple[int, int] | None = None,
) -> TagFacts:
    return TagFacts(
        name=name,
        category=category,
        must_try=must_try,
        notes=notes,
        hours=opening,
        best_window=best,
        viewpoints=("Langbiang",),
    )


FOOD, CAFE, SIGHT, OTHER = Category.FOOD, Category.CAFE, Category.ATTRACTION, Category.OTHER

TAG_CASES: list[tuple[TagFacts, list[str]]] = [
    (facts("Bún Bò Cô Bê"), ["breakfast"]),
    (facts("Quán 79", must_try=("bánh căn trứng",)), ["breakfast"]),
    (facts("Pho Bang"), ["breakfast"]),
    (facts("Bánh cuốn Cô Hoá", notes=("Ăn sáng (9h22)",)), ["breakfast"]),
    (facts("Lẩu Gà Lá É Tao Ngộ"), ["dinner"]),
    (facts("Tiệm Nướng Trong Rừng", must_try=("cơm lam", "lẩu")), ["lunch", "dinner"]),
    (facts("Bánh tráng nướng Dì Đinh"), ["snack"]),
    (facts("Chè Thái Lan"), ["snack", "dessert"]),
    (facts("Quán Thành Long", opening=hours("06:00", "11:00")), ["breakfast"]),
    (facts("Phở Đêm", opening=hours("17:00", "02:00")), ["dinner", "late_night"]),
    (facts("Cơm Niêu", opening=hours("10:00", "21:00")), ["lunch"]),
    (facts("Quán Bà Tư", opening=hours("10:00", "21:00")), ["lunch", "dinner"]),
    (facts("Xôi Gà", notes=("meal: dinner",)), ["breakfast", "dinner"]),
    (facts("Quán X", best=(17 * 60, 20 * 60)), ["dinner"]),
    (
        facts("Cà Phê Mê Linh", CAFE, notes=("Quán view đồi thông, ngắm hoàng hôn",)),
        ["view", "sunset"],
    ),
    (facts("Tiệm Bánh Cối Xay Gió", CAFE, must_try=("bánh su kem",)), ["dessert"]),
    (facts("Cà phê muối", CAFE, must_try=("cà phê kem muối",)), []),
    (facts("Cafe Workspace Dalat", CAFE, notes=("có ổ cắm, wifi mạnh",)), ["work_friendly"]),
    (
        facts("Đồi Đa Phú", SIGHT, notes=("Săn mây spot",)),
        ["view", "sunrise", "cloud_hunting", "nature"],
    ),
    (facts("Đường Nguyễn Khuyến", SIGHT, notes=("Check-in miễn phí",)), ["photo_spot"]),
    (facts("Vườn dâu tây Bích Câu", SIGHT), ["garden"]),
    (facts("Dinh Bảo Đại", SIGHT), ["history"]),
    (facts("Bảo tàng Hồ Chí Minh", SIGHT), ["history"]),
    (facts("Đỉnh Langbiang", SIGHT), ["view", "nature"]),
    (facts("Công viên Yersin", SIGHT), ["family"]),
    (facts("Chợ Đêm Đà Lạt", OTHER), ["shopping"]),
    (
        facts("Little Fish Homestay", OTHER, notes=("Chỗ ở (slide 'Homestay cần lưu')",)),
        ["service"],
    ),
    (facts("Gần quán bún", CAFE), []),
    (facts("Pink Valley", SIGHT), []),
]


@pytest.mark.parametrize(("case", "expected"), TAG_CASES, ids=[c.name for c, _ in TAG_CASES])
def test_tag_table(case: TagFacts, expected: list[str]) -> None:
    assert derive_tags(case) == expected


def test_vocabulary_is_complete() -> None:
    assert len(TAGS) == len(set(TAGS)) == 18
    for tag in TAGS:
        assert tag == tag.lower() and " " not in tag
        assert TAG_WINDOWS[tag] and TAG_WHY[tag]
        assert tag in BEST_TIME_ORDER
        assert EN.tags[tag] and VI.tags[tag]
    assert set(EN.tags) == set(VI.tags) == set(TAGS)


def test_meal_hint_wins_the_best_time() -> None:
    hits = tag_hits(facts("Xôi Gà", notes=("meal: dinner",)))
    assert (hits[0].tag, hits[0].reason) == ("dinner", "meal: dinner")
    late = tag_hits(facts("Quán Y", notes=("Gợi ý của bạn; meal: late night",)))
    assert [h.tag for h in late] == ["late_night"]
    brunch = derive_tags(facts("Quán Z", notes=("meal: brunch (slide 3)",)))
    assert brunch == ["breakfast", "lunch"]


def test_hours_drop_impossible_meal_tags() -> None:
    hits = tag_hits(facts("Phở Đêm", opening=hours("17:00", "02:00")))
    assert hits[0].tag == "dinner" and hits[0].reason == "opens 17:00"
    assert "breakfast" not in [h.tag for h in hits]


def test_evidence_best_time_feeds_tags_but_heuristic_does_not() -> None:
    place = Place(
        id="osm:node/1", name="Quán X", category=FOOD, address="a", lat=0, lng=0, confidence=0.9
    )
    evidence = BestTime(hours="07:00–09:00", source="comments")
    heuristic = BestTime(hours="07:00–09:00", source="heuristic")
    assert derive_tags(TagFacts.of(place, evidence)) == ["breakfast"]
    assert derive_tags(TagFacts.of(place, heuristic)) == []


def place_facts(case: TagFacts, *texts: SourceText) -> PlaceFacts:
    return PlaceFacts(
        name=case.name,
        category=case.category,
        must_try=case.must_try,
        hours=case.hours,
        texts=texts,
        profile=DA_LAT,
        tags=tuple(tag_hits(case)),
    )


@pytest.mark.parametrize("category", list(Category))
def test_every_place_gets_a_best_time_window(category: Category) -> None:
    result = suggest(place_facts(facts("Nơi Không Tên", category)))
    assert result is not None and result.hours
    assert result.source == "heuristic"
    assert result.why and result.why_vi


def test_tag_heuristic_names_the_tag_and_explains_why() -> None:
    result = suggest(place_facts(facts("Bún Bò Cô Bê", opening=hours("06:30", "15:00"))))
    assert result is not None
    assert result.hours == "06:30–09:00"
    assert result.evidence is not None and result.evidence.startswith('tag rule: breakfast ("bún"')
    assert result.why == "breakfast place, early morning, opens 06:30"
    assert result.why_vi == "quán ăn sáng, đi sớm, mở cửa 06:30"
    assert result.months == "Nov–Apr"

    cafe = suggest(place_facts(facts("Cà Phê Đồi", CAFE)))
    assert cafe is not None
    assert (cafe.hours, cafe.why) == ("15:00–17:30", "view café, afternoon light")
    clouds = suggest(place_facts(facts("Đồi Thiên Phúc Đức", SIGHT, notes=("săn mây",))))
    assert clouds is not None
    assert (clouds.hours, clouds.why) == ("04:30–07:00", "cloud hunting at dawn")


def test_heuristic_window_is_clamped_and_alternative_picked_by_hours() -> None:
    view = suggest(place_facts(facts("Cà Phê Đồi", CAFE, opening=hours("07:00", "16:00"))))
    assert view is not None
    assert view.hours == "15:00–16:00"
    assert view.evidence is not None and "adjusted to opening hours" in view.evidence
    sweets = facts("Tiệm Bánh", CAFE, must_try=("bánh flan",), opening=hours("17:00", "22:00"))
    dessert = suggest(place_facts(sweets))
    assert dessert is not None
    assert dessert.hours == "18:00–21:00"


def test_text_evidence_explains_itself_and_keeps_tag_hours_as_filler() -> None:
    evening = suggest(place_facts(facts("Ốc Đào"), SourceText("comments", "tối mới ngon")))
    assert evening is not None
    assert evening.source == "comments"
    assert evening.why == "comments: evening"
    assert evening.why_vi == "bình luận: buổi tối"

    weekend = suggest(
        place_facts(facts("Bún Bò Cô Bê"), SourceText("caption", "Chỉ bán cuối tuần thôi"))
    )
    assert weekend is not None
    assert (weekend.source, weekend.days, weekend.hours) == ("caption", "weekend", "06:30–09:00")
    assert weekend.why == "caption: weekends; breakfast place, early morning"


def test_enrich_sets_tags_and_why_and_fingerprints_tags() -> None:
    base = Place(
        id="osm:node/1",
        name="Quán Bà Tư",
        category=Category.FOOD,
        address="a",
        lat=0,
        lng=0,
        confidence=0.9,
    )
    plain = enrich.with_best_time(base, [], DA_LAT)
    assert plain.tags == []
    assert plain.best_time is not None and plain.best_time.why == "food place, dinner time"

    hinted = enrich.with_best_time(base.model_copy(update={"notes": ["meal: breakfast"]}), [])
    assert hinted.tags == ["breakfast"]
    assert hinted.best_time is not None
    assert hinted.best_time.hours == "06:30–09:00"
    assert hinted.best_time.inputs_hash != plain.best_time.inputs_hash
    assert enrich.with_best_time(hinted, []) == hinted


def test_labels_for_tags_and_why() -> None:
    assert tags_text(EN, ["breakfast", "late_night"]) == "breakfast · late night"
    assert tags_text(VI, ["breakfast", "view"]) == "ăn sáng · view đẹp"
    assert tags_text(EN, []) == ""
    best = BestTime(
        hours="06:30–09:00",
        source="heuristic",
        why="breakfast place, early morning, opens 06:30",
        why_vi="quán ăn sáng, đi sớm, mở cửa 06:30",
    )
    assert best_time_why(EN, best) == "breakfast place, early morning, opens 6:30 AM"
    assert best_time_why(VI, best) == "quán ăn sáng, đi sớm, mở cửa 6:30"
    assert best_time_why(VI, best.model_copy(update={"why_vi": None})).startswith("breakfast")
    assert best_time_why(EN, None) == ""
    assert EN.map_ui["tags"] == EN.tags_title and VI.map_ui["best_time_why"] == VI.best_time_why


def test_excel_shows_tags_why_and_estimate(tmp_path: Path) -> None:
    data = build_dataset()
    why = {"why": "breakfast place, early morning", "why_vi": "quán ăn sáng, đi sớm"}
    data.places = [
        p.model_copy(
            update={
                "tags": ["breakfast", "view"],
                "best_time": p.best_time.model_copy(update=why) if p.best_time else None,
            }
        )
        if p.id == "gp:pho1"
        else p
        for p in data.places
    ]
    en = load_workbook(write_trip(data, EN, tmp_path / "en.xlsx"))
    vi = load_workbook(write_trip(data, VI, tmp_path / "vi.xlsx"))

    assert en["Dashboard"]["E5"].value == (
        "7:00 AM–9:00 PM (best: 6:30–9:00 AM, estimate)\nbreakfast · view"
    )
    assert (
        vi["Tổng Hợp"]["E5"].value
        == "7:00–21:00 (giờ đẹp: 6:30–9:00, ước tính)\năn sáng · view đẹp"
    )

    plan, vi_plan = en["Plan-to-go"], vi["Lịch Trình"]
    assert (plan["M6"].value, plan["N6"].value) == ("Best time", "Tags")
    assert vi_plan["N6"].value == "Nhãn"
    pho = next(
        r for r in plan.iter_rows(min_row=7) if str(r[3].value).endswith('"Phở Hòa Pasteur")')
    )
    assert pho[12].value == "6:30–9:00 AM (estimate)\nbreakfast place, early morning"
    assert pho[13].value == "breakfast · view"
    vi_pho = next(
        r for r in vi_plan.iter_rows(min_row=7) if str(r[3].value).endswith('"Phở Hòa Pasteur")')
    )
    assert vi_pho[12].value == "6:30–9:00 (ước tính)\nquán ăn sáng, đi sớm"
    assert vi_pho[13].value == "ăn sáng · view đẹp"
